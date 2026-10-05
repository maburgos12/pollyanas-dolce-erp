"""Procedencia global y revisión estrictamente de lectura, en PostgreSQL."""
from io import BytesIO, StringIO
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from core.models import AuditLog
from activos.models import Activo, OrdenMantenimiento
from activos.utils import bitacora_import


class ImportacionBitacoraTests(TestCase):
    def setUp(self):
        self.actor = get_user_model().objects.create_superuser('importador', password='test')
        self.otro = get_user_model().objects.create_superuser('otro_importador', password='test')
        self.activo = Activo.objects.create(nombre='Equipo existente')

    def archivo(self, costo='1250.75', fecha='2026-02-20', segundo=False):
        tail = ',2026-02-20,99' if segundo else ',,'
        return SimpleUploadedFile('historial.csv', (
            'nombre,marca,modelo,serie,fecha_1,costo_1,fecha_2,costo_2\n'
            f'Equipo archivo,Marca,Modelo,S1,{fecha},{costo}{tail}\n').encode())

    def decision(self, fila=2, slot=1, **campos):
        return dict(fila=fila, slot=slot, accion='crear', activo_id=self.activo.pk,
                    orden_id=None, motivo='Revisado contra archivo', evidencia='Documento fuente', **campos)

    def preview(self, archivo):
        self.assertTrue(hasattr(bitacora_import, 'preview_bitacora'), 'Falta vista previa sin escrituras')
        return bitacora_import.preview_bitacora(archivo)

    def test_preview_no_dml_audit_signals_or_sequence_advance(self):
        from django.test.utils import CaptureQueriesContext
        with connection.cursor() as cursor:
            cursor.execute("SELECT last_value, is_called FROM activos_activo_id_seq")
            antes = cursor.fetchone()
        audits = AuditLog.objects.count()
        with CaptureQueriesContext(connection) as queries, patch('activos.utils.bitacora_import._upsert_activo', create=True, side_effect=AssertionError('No upsert')):
            resultado = self.preview(self.archivo())
        self.assertEqual(resultado['servicios'][0]['costo'], '1250.75')
        self.assertFalse(any(q['sql'].lstrip().split()[0].upper() in {'INSERT','UPDATE','DELETE'} or 'nextval' in q['sql'].lower() for q in queries))
        self.assertEqual(AuditLog.objects.count(), audits)
        self.assertEqual(Activo.objects.count(), 1)
        with connection.cursor() as cursor:
            cursor.execute("SELECT last_value, is_called FROM activos_activo_id_seq")
            self.assertEqual(cursor.fetchone(), antes)

    def test_template_xlsx_round_trip_and_real_sheet(self):
        self.client.force_login(self.actor)
        response = self.client.get(reverse('activos:activos'), {'export':'template_bitacora_xlsx'})
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(response.content)); ws = wb.active
        ws.append(['Equipo plantilla','Marca','Modelo','S2','2026-02-20',100,'',''])
        stream = BytesIO(); wb.save(stream)
        result = self.preview(SimpleUploadedFile('plantilla.xlsx', stream.getvalue()))
        fila = next(s for s in result['servicios'] if s['nombre'] == 'Equipo plantilla')
        self.assertEqual(fila['costo'], '100.00')
        self.assertEqual(result['sheet_name'], ws.title)

    def test_csv_physical_start_lines_and_decimal_comma(self):
        archivo = SimpleUploadedFile('lineas.csv', ('\ufeffnombre;"marca\n";modelo;serie;fecha_1;costo_1;fecha_2;costo_2\n'
            '"Equipo\nmultilinea";M;X;S;2026-02-20;1.250,75;;\nOtro;M;X;S;2026-02-21;100;;').encode())
        result = self.preview(archivo)
        self.assertEqual([s['fila'] for s in result['servicios']], [3,5])
        self.assertEqual(result['servicios'][0]['costo'], '1250.75')

    def test_unknown_cost_never_becomes_zero(self):
        for costo in ('', 'texto', 'NaN', 'Infinity', '-Infinity', '1e9999'):
            with self.subTest(costo=costo):
                self.assertIsNone(self.preview(self.archivo(costo))['servicios'][0]['costo'])

    def aplicar(self, archivo=None, decisiones=None, actor=None):
        from activos import services_importacion as servicio
        return servicio.aplicar_importacion(usuario=actor or self.actor, archivo=archivo or self.archivo(),
                                            decisiones=decisiones or [self.decision()], confirmado=True)

    def test_create_explicit_existing_asset_and_two_slots(self):
        from activos import services_importacion as servicio
        result = servicio.aplicar_importacion(usuario=self.actor, archivo=self.archivo(segundo=True),
            decisiones=[self.decision(), self.decision(slot=2)], confirmado=True)
        self.assertEqual(Activo.objects.count(), 1)
        self.assertEqual(OrdenMantenimiento.objects.count(), 2)
        self.assertNotEqual(result[0]['orden_id'], result[1]['orden_id'])
        orden = OrdenMantenimiento.objects.get(pk=result[0]['orden_id'])
        self.assertEqual(orden.costo_otros, Decimal('1250.75'))
        self.assertEqual(orden.creado_por, self.actor)

    def test_global_replay_cross_actor_preserves_author_and_audit(self):
        from activos.models import OrigenImportacionBitacora
        first = self.aplicar()
        audits = AuditLog.objects.count()
        second = self.aplicar(actor=self.otro)
        self.assertEqual(first[0]['orden_id'], second[0]['orden_id'])
        self.assertTrue(second[0]['reintento'])
        self.assertEqual(AuditLog.objects.count(), audits)
        self.assertEqual(OrigenImportacionBitacora.objects.get().autor_original_id, self.actor.pk)

    def test_conflict_on_changed_decision_or_evidence(self):
        from mantenimiento.services_capturas_equipos import CapturaEquipoError
        self.aplicar()
        for cambio in ({'evidencia':'otra'}, {'motivo':'otro'}, {'accion':'vincular','orden_id':OrdenMantenimiento.objects.get().pk}):
            decision = self.decision(); decision.update(cambio)
            with self.assertRaises(CapturaEquipoError) as exc:
                self.aplicar(decisiones=[decision])
            self.assertEqual(exc.exception.status_code, 409)
        self.assertEqual(OrdenMantenimiento.objects.count(), 1)

    def test_tombstone_after_current_permissions_and_no_recreation(self):
        from mantenimiento.services_capturas_equipos import CapturaEquipoError
        self.aplicar()
        OrdenMantenimiento.objects.get().delete()
        with self.assertRaises(CapturaEquipoError) as exc:
            self.aplicar()
        self.assertEqual(exc.exception.status_code, 410)
        self.actor.is_active = False; self.actor.save(update_fields=['is_active'])
        with self.assertRaises(PermissionDenied):
            self.aplicar()
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)

    def test_unknown_cost_can_link_without_mutating_existing_order(self):
        orden = OrdenMantenimiento.objects.create(activo_ref=self.activo, costo_otros=Decimal('42'))
        decision = self.decision(); decision.update(accion='vincular', orden_id=orden.pk)
        self.aplicar(archivo=self.archivo(''), decisiones=[decision])
        orden.refresh_from_db(); self.assertEqual(orden.costo_otros, Decimal('42'))
        self.assertEqual(OrdenMantenimiento.objects.count(), 1)

    def test_unknown_cost_or_invalid_date_cannot_create(self):
        from mantenimiento.services_capturas_equipos import CapturaEquipoError
        for archivo in (self.archivo(''), self.archivo('NaN'), self.archivo(fecha='no fecha')):
            with self.assertRaises(CapturaEquipoError) as exc:
                self.aplicar(archivo=archivo)
            self.assertEqual(exc.exception.status_code, 400)
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)

    def test_order_of_other_equipment_cannot_be_linked(self):
        otro = Activo.objects.create(nombre='Otro equipo')
        orden = OrdenMantenimiento.objects.create(activo_ref=otro)
        decision = self.decision(); decision.update(accion='vincular', orden_id=orden.pk)
        with self.assertRaises(PermissionDenied):
            self.aplicar(decisiones=[decision])

    def test_rollback_second_failure_removes_all_work_origins_and_audits(self):
        from activos.models import OrigenImportacionBitacora
        from activos import services_importacion as servicio
        before = AuditLog.objects.count()
        original = servicio.capturar_equipo_autorizado
        calls = []
        def falla(**kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError('fallo intermedio')
            return original(**kwargs)
        with patch.object(servicio, 'capturar_equipo_autorizado', side_effect=falla):
            with self.assertRaises(RuntimeError):
                self.aplicar(archivo=self.archivo(segundo=True), decisiones=[self.decision(),self.decision(slot=2)])
        self.assertEqual(OrigenImportacionBitacora.objects.count(), 0)
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)
        self.assertEqual(AuditLog.objects.count(), before)

    def test_web_three_steps_readonly_then_confirm_and_cli_replay(self):
        from django.core.management import call_command
        from django.test.utils import CaptureQueriesContext
        from tempfile import TemporaryDirectory
        from pathlib import Path
        import json
        self.client.force_login(self.actor)
        url = reverse('activos:activos') + '?q=Equipo'
        with CaptureQueriesContext(connection) as queries:
            response = self.client.post(url, {'action':'import_bitacora', 'archivo_bitacora':self.archivo()}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 200)
        self.assertFalse(any(q['sql'].lstrip().split()[0].upper() in {'INSERT','UPDATE','DELETE'} or 'nextval' in q['sql'].lower() for q in queries))
        payload = response.json()
        self.assertEqual(payload['target'], '#importacion-bitacora')
        token = payload['archivo_token']
        decision = self.decision()
        campos = {f'{k}_2_1':v for k,v in decision.items() if k not in ('fila','slot') and v is not None}
        review = self.client.post(url, dict(action='import_bitacora', fase='review', archivo_token=token, **campos), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(review.status_code, 200)
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)
        confirm = self.client.post(url, dict(action='import_bitacora', fase='confirm', revision_token=review.json()['revision_token'], confirmado='1'), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(confirm.status_code, 200)
        self.assertEqual(OrdenMantenimiento.objects.count(), 1)
        audits = AuditLog.objects.count()
        with TemporaryDirectory() as folder:
            archivo = Path(folder)/'historial.csv'; archivo.write_bytes(self.archivo().read())
            revisado = Path(folder)/'revision.json'
            revisado.write_text(json.dumps({'archivo_sha256':self.preview(self.archivo())['sha256'], 'hoja':'CSV','revisado':True,'decisiones':[decision]}))
            call_command('importar_activos_bitacora', str(archivo), apply=True, actor=str(self.otro.pk), decisiones=str(revisado), stdout=StringIO())
        self.assertEqual(OrdenMantenimiento.objects.count(), 1)
        self.assertEqual(AuditLog.objects.count(), audits)

    def test_web_legacy_apply_only_previews_no_master_creation(self):
        self.client.force_login(self.actor)
        response = self.client.post(reverse('activos:activos'), {'action':'import_bitacora','archivo_bitacora':self.archivo()})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Revisar decisiones')
        self.assertEqual(Activo.objects.count(), 1)
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)
        self.assertEqual(AuditLog.objects.filter(model='activos.BitacoraImport').count(), 0)

    def test_cli_apply_cannot_bypass_explicit_actor_and_review(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from django.core.management import call_command, CommandError
        with TemporaryDirectory() as folder:
            archivo = Path(folder)/'historial.csv'; archivo.write_bytes(self.archivo().read())
            with self.assertRaises(CommandError):
                call_command('importar_activos_bitacora',str(archivo),apply=True, stdout=StringIO())
        self.assertEqual(Activo.objects.count(), 1)
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)

    def scoped_actor(self):
        from core.models import Sucursal, UserModuleAccess, UserProfile
        branch = Sucursal.objects.create(codigo='IMPORT-A',nombre='Sucursal A')
        other = Sucursal.objects.create(codigo='IMPORT-B',nombre='Sucursal B')
        actor = get_user_model().objects.create_user('captura_sucursal')
        UserProfile.objects.update_or_create(user=actor, defaults={'sucursal':branch})
        UserModuleAccess.objects.create(user=actor,module='inventario',access='manage')
        UserModuleAccess.objects.create(user=actor,module='mantenimiento.app',access='manage')
        self.activo.sucursal = branch; self.activo.save(update_fields=['sucursal'])
        return actor, branch, other

    def test_native_write_scope_uses_existing_capture_permission(self):
        from mantenimiento.services_access import can_view_costs
        actor, branch, other = self.scoped_actor()
        self.aplicar(actor=actor)
        self.assertEqual(OrdenMantenimiento.objects.count(), 1)

    def test_fresh_revocation_even_same_actor_cached_permissions(self):
        from core.access import can_manage_inventario
        from core.models import UserModuleAccess
        actor, _, _ = self.scoped_actor()
        self.assertTrue(can_manage_inventario(actor))
        self.aplicar(actor=actor)
        before = AuditLog.objects.count()
        UserModuleAccess.objects.filter(user=actor,module='mantenimiento.app').update(access='none')
        for cambio in ({}, {'motivo':'conflicto no visible'}):
            decision = self.decision(); decision.update(cambio)
            with self.assertRaises(PermissionDenied):
                self.aplicar(actor=actor, decisiones=[decision])
        self.assertEqual(AuditLog.objects.count(), before)
        OrdenMantenimiento.objects.get().delete()
        with self.assertRaises(PermissionDenied):
            self.aplicar(actor=actor)

    def test_other_branch_and_inactive_equipment_are_forbidden(self):
        actor, _, other = self.scoped_actor()
        self.activo.sucursal = other; self.activo.save(update_fields=['sucursal'])
        with self.assertRaises(PermissionDenied):
            self.aplicar(actor=actor)
        self.activo.activo = False; self.activo.save(update_fields=['activo'])
        with self.assertRaises(PermissionDenied):
            self.aplicar()
        self.assertEqual(OrdenMantenimiento.objects.count(), 0)

    def test_linked_order_tombstone_returns_410(self):
        from mantenimiento.services_capturas_equipos import CapturaEquipoError
        orden = OrdenMantenimiento.objects.create(activo_ref=self.activo)
        decision = self.decision(); decision.update(accion='vincular',orden_id=orden.pk)
        self.aplicar(decisiones=[decision])
        orden.delete()
        with self.assertRaises(CapturaEquipoError) as exc:
            self.aplicar(decisiones=[decision])
        self.assertEqual(exc.exception.status_code, 410)

    def test_exact_file_identity_and_corrected_bytes_need_new_review(self):
        from activos.models import OrigenImportacionBitacora
        first = self.aplicar()
        archivo = self.archivo('1251.75'); archivo.name='mismo_nombre.csv'
        second = self.aplicar(archivo=archivo)
        self.assertNotEqual(first[0]['orden_id'],second[0]['orden_id'])
        self.assertEqual(OrigenImportacionBitacora.objects.count(),2)

    def test_discarded_unselected_and_ids_outside_limits_do_not_write(self):
        from mantenimiento.services_capturas_equipos import CapturaEquipoError
        from activos.services_importacion import aplicar_importacion
        for accion in ('descartar',''):
            decision = self.decision(); decision['accion']=accion
            self.assertEqual(self.aplicar(decisiones=[decision]), [])
        for value in ('9'*100, -1, 'NaN'):
            decision = self.decision(); decision['activo_id']=value
            with self.assertRaises(CapturaEquipoError) as exc:
                self.aplicar(decisiones=[decision])
            self.assertEqual(exc.exception.status_code,400)
        self.assertEqual(OrdenMantenimiento.objects.count(),0)

    def test_source_exclusions_legacy_sheets_and_formula_unknown(self):
        wb = Workbook(); ws=wb.active; ws.title='Historia real'
        ws.append(['','LOGISTICA','','','','','','',''])
        ws.append(['','Camión','X','Z','S','2026-02-20',100,'',''])
        ws.append(['','MATRIZ','','','','','','',''])
        ws.append(['','Obra','X','Z','INSTALACION','2026-02-20',100,'',''])
        ws.append(['','CORTINA METALICA','X','Z','INSTALACION','=DATE(2026,2,20)','=10+10','',''])
        otro = wb.create_sheet('Otra hoja'); otro.append(['nombre','marca','modelo','serie','fecha_1','costo_1','fecha_2','costo_2']); otro.append(['Otro','X','Y','S','2026-02-20',50,'',''])
        raw = BytesIO(); wb.save(raw)
        upload = SimpleUploadedFile('legado.xlsx',raw.getvalue())
        preview = self.preview(upload)
        self.assertEqual([(s['nombre'],s['fila'],s['costo']) for s in preview['servicios']], [('CORTINA METALICA',5,None)])
        self.assertEqual(bitacora_import.preview_bitacora(upload,sheet_name='Otra hoja')['servicios'][0]['nombre'],'Otro')

    def test_bounds_huge_finite_cost_and_rounding_overflow_unknown(self):
        for costo in ('9999999999999999.995','99999999999999999', '1e16'):
            self.assertIsNone(self.preview(self.archivo(costo))['servicios'][0]['costo'])
        with self.assertRaises(ValueError):
            self.preview(SimpleUploadedFile('big.csv',b'x'*(2*1024*1024+1)))
        with self.assertRaises(ValueError):
            self.preview(SimpleUploadedFile('rows.csv',self.archivo().read()*501))

    def test_review_itself_has_no_dml_or_sequence_calls(self):
        from activos.services_importacion import revisar_importacion
        from django.test.utils import CaptureQueriesContext
        with CaptureQueriesContext(connection) as queries:
            actor, preview, rows = revisar_importacion(usuario=self.actor,archivo=self.archivo(),decisiones=[self.decision()])
        self.assertEqual(len(rows),1)
        self.assertFalse(any(q['sql'].lstrip().split()[0].upper() in {'INSERT','UPDATE','DELETE'} or 'nextval' in q['sql'].lower() for q in queries))

    def test_native_paging_keeps_draft_and_review_back_keeps_selections(self):
        self.client.force_login(self.actor)
        contenido = ('nombre,marca,modelo,serie,fecha_1,costo_1,fecha_2,costo_2\n' +
            ''.join(f'Equipo {i},X,Y,S,2026-02-20,100,,\n' for i in range(105))).encode()
        url=reverse('activos:activos')
        preview=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=SimpleUploadedFile('muchos.csv',contenido)),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        campos=dict(accion_2_1='crear',activo_id_2_1=self.activo.pk,motivo_2_1='Primera selección',evidencia_2_1='Fuente')
        page=self.client.post(url,dict(action='import_bitacora',fase='review',mover_pagina='1',archivo_token=preview['archivo_token'],**campos),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(page.status_code,200)
        self.assertIn('Fila 102',page.json()['html'])
        review=self.client.post(url,dict(action='import_bitacora',fase='review',pagina='1',archivo_token=page.json()['archivo_token'],accion_102_1='crear',activo_id_102_1=self.activo.pk,motivo_102_1='Segunda selección',evidencia_102_1='Fuente'),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(review.status_code,200)
        self.assertIn('Primera selección',review.json()['html'])
        self.assertIn('Segunda selección',review.json()['html'])
        back=self.client.post(url,dict(action='import_bitacora',fase='edit',archivo_token=review.json()['revision_token']),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(back.status_code,200)
        self.assertIn('Primera selección',back.json()['html'])
        confirm=self.client.post(url,dict(action='import_bitacora',fase='confirm',revision_token=review.json()['revision_token'],confirmado='1'),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(confirm.status_code,200)
        self.assertEqual(OrdenMantenimiento.objects.count(),2)

    def test_confirmation_error_preserves_file_and_review_decisions(self):
        self.client.force_login(self.actor); url=reverse('activos:activos')
        preview=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        review=self.client.post(url,dict(action='import_bitacora',fase='review',archivo_token=preview['archivo_token'],accion_2_1='crear',activo_id_2_1=self.activo.pk,motivo_2_1='Conservar motivo',evidencia_2_1='Conservar evidencia'),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        failed=self.client.post(url,dict(action='import_bitacora',fase='confirm',revision_token=review['revision_token']),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(failed.status_code,400)
        self.assertIn('Conservar motivo',failed.json()['html'])
        self.assertIn('Conservar evidencia',failed.json()['html'])
        self.assertTrue(failed.json()['archivo_token'])
        self.assertEqual(failed.json()['revision_token'],review['revision_token'])
        self.assertEqual(OrdenMantenimiento.objects.count(),0)


    def test_preview_token_cannot_confirm_unreviewed_draft(self):
        self.client.force_login(self.actor); url=reverse('activos:activos')
        preview=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        page=self.client.post(url,dict(action='import_bitacora',fase='review',mover_pagina='0',archivo_token=preview['archivo_token'],accion_2_1='crear',activo_id_2_1=self.activo.pk,motivo_2_1='Borrador',evidencia_2_1='Fuente'),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        response=self.client.post(url,dict(action='import_bitacora',fase='confirm',revision_token=page['archivo_token'],confirmado='1'),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,400)
        self.assertEqual(OrdenMantenimiento.objects.count(),0)

    def test_expired_confirmation_keeps_file_and_selections(self):
        self.client.force_login(self.actor); url=reverse('activos:activos')
        preview=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        review=self.client.post(url,dict(action='import_bitacora',fase='review',archivo_token=preview['archivo_token'],accion_2_1='crear',activo_id_2_1=self.activo.pk,motivo_2_1='Revisión conservada',evidencia_2_1='Fuente'),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        import time
        with patch('django.core.signing.time.time',return_value=time.time()+25*60*60):
            response=self.client.post(url,dict(action='import_bitacora',fase='confirm',revision_token=review['revision_token'],confirmado='1'))
        self.assertEqual(response.status_code,409)
        self.assertContains(response,'Revisión conservada',status_code=409)
        self.assertContains(response,'name="archivo_token"',status_code=409)
        self.assertEqual(OrdenMantenimiento.objects.count(),0)

    def test_final_token_limit_and_non_object_cli_review(self):
        from activos.views_importacion import _token
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from django.core.management import call_command, CommandError
        import os, base64
        with self.assertRaises(ValueError):
            _token({'bytes':base64.b64encode(os.urandom(1600000)).decode()})
        with TemporaryDirectory() as folder:
            archivo=Path(folder)/'historial.csv'; archivo.write_bytes(self.archivo().read())
            decisiones=Path(folder)/'revision.json'
            for raw in ('[]','null','42','"texto"'):
                decisiones.write_text(raw)
                with self.assertRaises(CommandError):
                    call_command('importar_activos_bitacora',str(archivo),apply=True,actor=str(self.actor.pk),decisiones=str(decisiones),stdout=StringIO())

    def test_malformed_file_missing_file_and_invalid_token_are_400(self):
        self.client.force_login(self.actor); url=reverse('activos:activos')
        for data in (dict(archivo_bitacora=SimpleUploadedFile('mal.xlsx',b'no zip')),
                     {},dict(fase='confirm',revision_token='falso',confirmado='1')):
            response=self.client.post(url,dict(action='import_bitacora',**data),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
            self.assertEqual(response.status_code,400)
        self.assertEqual(AuditLog.objects.filter(model='activos.BitacoraImport').count(),0)


    def test_active_session_revoked_scope_returns_json_403(self):
        actor, branch, other = self.scoped_actor()
        self.client.force_login(actor); url=reverse('activos:activos')
        preview=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.activo.sucursal=other; self.activo.save(update_fields=['sucursal'])
        response=self.client.post(url,dict(action='import_bitacora',fase='review',archivo_token=preview['archivo_token'],accion_2_1='crear',activo_id_2_1=self.activo.pk,motivo_2_1='Revisado',evidencia_2_1='Fuente'),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,403)
        self.assertTrue(response['Content-Type'].startswith('application/json'))
        self.assertFalse(response.json()['ok'])
        self.assertIn('toast',response.json())
        self.assertNotIn('html',response.json())
        from core.models import UserModuleAccess
        UserModuleAccess.objects.filter(user=actor,module='inventario').update(access='none')
        response=self.client.post(url,dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,403)
        self.assertTrue(response['Content-Type'].startswith('application/json'))

    def test_inactive_session_keeps_native_login_redirect_and_snapshot_contract(self):
        self.client.force_login(self.actor)
        preview=self.client.post(reverse('activos:activos'),dict(action='import_bitacora',archivo_bitacora=self.archivo()),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertIn('data-capture-snapshot="true"',preview.json()['html'])
        self.actor.is_active=False; self.actor.save(update_fields=['is_active'])
        response=self.client.post(reverse('activos:activos'),dict(action='import_bitacora',archivo_token=preview.json()['archivo_token']),HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,302)
        self.assertIn('login',response['Location'])


    def test_cli_apply_uses_reviewed_bytes_when_source_path_changes(self):
        from activos.management.commands import importar_activos_bitacora as command
        from activos.models import OrigenImportacionBitacora
        from django.core.management import call_command
        from pathlib import Path
        from tempfile import TemporaryDirectory
        import json
        import hashlib
        original_bytes = self.archivo('100').read()
        changed_bytes = self.archivo('200').read()
        original_sha = hashlib.sha256(original_bytes).hexdigest()
        real_preview = command.preview_bitacora
        with TemporaryDirectory() as folder:
            archivo = Path(folder) / 'historial.csv'
            archivo.write_bytes(original_bytes)
            decisiones = Path(folder) / 'revision.json'
            decisiones.write_text(json.dumps(dict(archivo_sha256=original_sha, hoja='CSV',
                revisado=True, decisiones=[self.decision()])))
            def reemplazar_despues_de_preview(source, **kwargs):
                preview = real_preview(source, **kwargs)
                self.assertEqual(preview['sha256'], original_sha)
                archivo.write_bytes(changed_bytes)
                return preview
            with patch.object(command, 'preview_bitacora', side_effect=reemplazar_despues_de_preview):
                call_command('importar_activos_bitacora', str(archivo), apply=True,
                    actor=str(self.actor.pk), decisiones=str(decisiones), stdout=StringIO())
            self.assertEqual(archivo.read_bytes(), changed_bytes)
        self.assertEqual(OrdenMantenimiento.objects.get().costo_otros, Decimal('100'))
        self.assertEqual(OrigenImportacionBitacora.objects.get().archivo_sha256, original_sha)


from django.test import TransactionTestCase
from django.db import close_old_connections, transaction
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event


class ImportacionConcurrenteTests(TransactionTestCase):
    """Carreras reales en PostgreSQL, sin SQLite ni mock de la transacción."""
    def setUp(self):
        self.actor=get_user_model().objects.create_superuser('concurrente',password='test')
        self.otro=get_user_model().objects.create_superuser('concurrente_otro',password='test')
        self.activo=Activo.objects.create(nombre='Equipo concurrente')
        self.decision=dict(fila=2,slot=1,accion='crear',activo_id=self.activo.pk,orden_id=None,motivo='Revisado',evidencia='Fuente')
        self.bytes=b'nombre,marca,modelo,serie,fecha_1,costo_1,fecha_2,costo_2\nEquipo,X,Y,S,2026-02-20,100,,\n'

    def worker(self, user_id, decision=None):
        from activos.services_importacion import aplicar_importacion
        close_old_connections()
        try:
            actor=get_user_model().objects.get(pk=user_id)
            return aplicar_importacion(usuario=actor, archivo=SimpleUploadedFile('race.csv',self.bytes),
                decisiones=[decision or self.decision], confirmado=True)
        finally:
            connection.close()

    def test_concurrent_cross_actor_global_unique(self):
        from activos import services_importacion as servicio
        from activos.models import OrigenImportacionBitacora
        barrier=Barrier(2); real=servicio.revisar_importacion
        def revisar(**kwargs):
            result=real(**kwargs); barrier.wait(timeout=10); return result
        with patch.object(servicio,'revisar_importacion',side_effect=revisar), ThreadPoolExecutor(max_workers=2) as pool:
            futures=[pool.submit(self.worker,u.pk) for u in (self.actor,self.otro)]
            resultados=[f.result(timeout=20) for f in futures]
        self.assertEqual(resultados[0][0]['orden_id'],resultados[1][0]['orden_id'])
        self.assertEqual(OrigenImportacionBitacora.objects.count(),1)
        self.assertEqual(OrdenMantenimiento.objects.count(),1)
        self.assertEqual(AuditLog.objects.filter(action='CREATE',model='activos.OrdenMantenimiento').count(),1)
        self.assertEqual(AuditLog.objects.filter(action='IMPORT',model='activos.BitacoraImport').count(),1)

    def test_permission_revoked_while_waiting_for_destination_lock(self):
        from activos import services_importacion as servicio
        from activos.models import OrigenImportacionBitacora
        reviewed=Event(); real=servicio.revisar_importacion
        def revisar(**kwargs):
            result=real(**kwargs); reviewed.set(); return result
        with patch.object(servicio,'revisar_importacion',side_effect=revisar), ThreadPoolExecutor(max_workers=1) as pool:
            with transaction.atomic():
                Activo.objects.select_for_update(no_key=True).get(pk=self.activo.pk)
                future=pool.submit(self.worker,self.actor.pk)
                self.assertTrue(reviewed.wait(timeout=10))
                get_user_model().objects.filter(pk=self.actor.pk).update(is_active=False)
            with self.assertRaises(PermissionDenied):
                future.result(timeout=20)
        self.assertEqual(OrdenMantenimiento.objects.count(),0)
        self.assertEqual(OrigenImportacionBitacora.objects.count(),0)

    def test_order_moved_to_other_equipment_during_lock_wait(self):
        from activos import services_importacion as servicio
        from activos.models import OrigenImportacionBitacora
        otro=Activo.objects.create(nombre='Destino ajeno')
        orden=OrdenMantenimiento.objects.create(activo_ref=self.activo)
        decision={**self.decision,'accion':'vincular','orden_id':orden.pk}
        reviewed=Event(); real=servicio.revisar_importacion
        def revisar(**kwargs):
            result=real(**kwargs); reviewed.set(); return result
        with patch.object(servicio,'revisar_importacion',side_effect=revisar), ThreadPoolExecutor(max_workers=1) as pool:
            with transaction.atomic():
                OrdenMantenimiento.objects.select_for_update(no_key=True).get(pk=orden.pk)
                future=pool.submit(self.worker,self.actor.pk,decision)
                self.assertTrue(reviewed.wait(timeout=10))
                OrdenMantenimiento.objects.filter(pk=orden.pk).update(activo_ref=otro)
            with self.assertRaises(PermissionDenied):
                future.result(timeout=20)
        self.assertEqual(OrigenImportacionBitacora.objects.count(),0)
        self.assertEqual(AuditLog.objects.filter(action='IMPORT',model='activos.BitacoraImport').count(),0)


from django.test import SimpleTestCase


class ImportacionSnapshotUITests(SimpleTestCase):
    def test_conflict_releases_only_readonly_import_snapshot_with_delegated_capture(self):
        import re
        import shutil
        import subprocess
        from pathlib import Path

        node = shutil.which('node')
        self.assertIsNotNone(node, 'El check del listener requiere Node.js disponible.')
        template = Path(__file__).parent / 'templates' / 'activos' / 'activos.html'
        script = re.search(r'{% block extra_js %}\s*<script>(.*?)</script>', template.read_text(), re.S).group(1)
        harness = r'''
const assert = require('node:assert/strict');
let onError;
global.document = {addEventListener: (name, callback, capture) => {
  assert.equal(name, 'erp:action-error');
  assert.equal(capture, true); // Shared CustomEvent does not bubble.
  onError = callback;
}};
'''
        checks = r'''
function form(fase, ownModule = true) {
  return {elements: {fase: {value: fase}}, _captureSnapshot: {motivo: 'B'},
    motivo: 'A', matches: selector => {
      assert.equal(selector, '#importacion-bitacora form[data-async-action]');
      return ownModule;
    }};
}
for (const fase of ['preview', 'review', 'page', 'edit']) {
  const readonly = form(fase);
  onError({target: readonly, detail: {statusCode: 409}});
  assert.equal(readonly._captureSnapshot, null);
  assert.equal(readonly.motivo, 'A'); // DOM keeps the operator correction.
}
for (const [fase, status, own] of [['confirm', 409, true], ['review', 0, true],
    ['review', 403, true], ['review', 409, false]]) {
  const stable = form(fase, own), original = stable._captureSnapshot;
  onError({target: stable, detail: {statusCode: status}});
  assert.equal(stable._captureSnapshot, original);
}
'''
        result = subprocess.run([node, '-e', harness + script + checks], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
