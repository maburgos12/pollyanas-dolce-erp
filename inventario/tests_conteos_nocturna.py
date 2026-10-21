from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from core.models import Sucursal, UserProfile
from pos_bridge.models import PointBranch, PointProduct, PointSyncJob, PointInventorySnapshot
from pos_bridge.services.product_count_units import COUNT_UNIT_KEY, catalog_count_unit
from pos_bridge.services.audit_stock_history_service import AuditStockHistoryService, AuditStockHistoryError
from inventario.conteos_conciliacion import MZ, cierre_dia
from inventario.conteos_nocturna import _adquirir_historial, _historial, _publicar, conciliar_conteos_nocturnos, huella_lecturas
from inventario.models_conteos import ConteoSucursal
from inventario.services_conteos import preparar_conteo, ejecutar_accion, lecturas_efectivas
from inventario.tests_conteos_conciliacion import movement


class NightlyCountTests(TestCase):
    def setUp(self):
        self.admin=get_user_model().objects.create_superuser('night-admin','admin@example.invalid','test-local')
        self.operator=get_user_model().objects.create_user('night-counter',password='test-local')
        self.branch=Sucursal.objects.create(codigo='NIGHT',nombre='Sucursal prueba nocturna')
        UserProfile.objects.update_or_create(user=self.operator,defaults={'sucursal':self.branch})
        self.point=PointBranch.objects.create(external_id='81',name='Sucursal prueba nocturna',erp_branch=self.branch)
        self.product=PointProduct.objects.create(external_id='100',sku='C100',name='Producto prueba',metadata={COUNT_UNIT_KEY:catalog_count_unit({'Codigo':'C100','Unidad':'PZA'})})
        today=timezone.now().astimezone(MZ).date()
        self.close=cierre_dia(today-timedelta(days=1))
        self.when=self.close-timedelta(hours=4)
        self.count=preparar_conteo(actor=self.admin,sucursal=self.branch,responsable=self.operator,
            fecha=today-timedelta(days=1),titulo='Conteo nocturno',items=[{'producto_id':self.product.pk}],request_id=uuid4())
        with patch('inventario.services_conteos.timezone.now',return_value=self.when-timedelta(minutes=5)):
            self.action('iniciar',{})
        self.line=self.count.lineas.get()
        self.action('enviar',{'lecturas':{str(self.line.pk):{'cantidad':'3','contado_en':self.when.isoformat()}},'observaciones':'Nota conservada'})
        self.job=PointSyncJob.objects.create(status='SUCCESS',job_type='inventory',
            started_at=self.close+timedelta(minutes=29),finished_at=self.close+timedelta(minutes=31))
        self.snapshot=PointInventorySnapshot.objects.create(branch=self.point,product=self.product,stock=3,
            sync_job=self.job,captured_at=self.close+timedelta(minutes=30),
            raw_payload={'headers':['Codigo','Cantidad','Unidad'],'row':['C100','3','PZA']})
        self.rows=[movement(101,self.when+timedelta(minutes=15),'VENTA','3','2'),
                   movement(102,self.when+timedelta(minutes=30),'ENTRADA POR TRANSFERENCIA','2','4'),
                   movement(103,self.close+timedelta(minutes=15),'VENTA','4','3')]
        # Point's naive Fecha contract is UTC, not the application timezone.
        for row in self.rows:
            row['Fecha']=row['Fecha'].removesuffix('+00:00')
        self.temp=TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings=SimpleNamespace(raw_exports_dir=Path(self.temp.name))
        self.fake=MagicMock()
        self.fake.get_stock_history.return_value=self.rows
        self.fake.__enter__.return_value=self.fake

    def action(self,name,payload,actor=None):
        self.count.refresh_from_db()
        result=ejecutar_accion(conteo_id=self.count.pk,actor=actor or self.operator,
            action=name,version=self.count.version,request_id=uuid4(),payload=payload)
        self.count.refresh_from_db()
        return result

    def acquire(self):
        return _adquirir_historial(self.fake,self.settings,self.point,self.product,self.count.fecha)

    def test_preserved_batch_is_used_once_without_touching_capture_or_reference(self):
        version=self.count.version
        readings=lecturas_efectivas(self.count)
        before=huella_lecturas(readings)
        with patch('inventario.conteos_nocturna.load_point_bridge_settings',return_value=self.settings), patch('inventario.conteos_nocturna.PointHttpSessionClient',return_value=self.fake):
            summary=conciliar_conteos_nocturnos()
        self.count.refresh_from_db()
        self.assertEqual(summary['calculados'],1)
        result=self.count.conciliacion['lineas'][str(self.line.pk)]
        self.assertEqual(Decimal(result['fisico_al_cierre']),4)
        self.assertEqual(Decimal(result['point_al_cierre']),4)
        self.assertEqual(result['posteriores_cierre'],[103])
        self.assertEqual(huella_lecturas(lecturas_efectivas(self.count)),before)
        self.assertEqual((self.count.version,self.count.estado,self.count.observaciones,self.count.referencia),(version,'ENVIADO','Nota conservada',{}))
        self.fake.login.assert_called_once()
        self.fake.get_stock_history.assert_called_once_with('100','81',movements=500)
        event=self.count.eventos.get(action='conciliacion_nocturna')
        self.assertIsNone(event.actor_id)
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            self.assertEqual(conciliar_conteos_nocturnos()['procesados'],0)
            client.assert_not_called()

    def test_missing_physical_time_stays_pending_without_point_contact(self):
        self.line.lecturas.update(contado_en=None)
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            summary=conciliar_conteos_nocturnos()
            client.assert_not_called()
        self.count.refresh_from_db()
        result=self.count.conciliacion['lineas'][str(self.line.pk)]
        self.assertIsNone(result['diferencia'])
        self.assertIn('Falta hora',result['motivo'])
        self.assertFalse(summary['reintentar'])

    def test_busy_point_is_pending_and_never_changes_the_physical_quantity(self):
        from contextlib import nullcontext
        with patch('inventario.conteos_nocturna.point_account_session_lock',return_value=nullcontext(False)), patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            result=conciliar_conteos_nocturnos()
            client.assert_not_called()
        self.count.refresh_from_db()
        self.assertTrue(result['reintentar'])
        self.assertIsNone(self.count.conciliacion['lineas'][str(self.line.pk)]['diferencia'])
        self.assertEqual(self.line.lecturas.get().cantidad,3)

    def test_task_retries_source_unavailability_and_schedule_is_two_am(self):
        from django.conf import settings
        from celery.exceptions import Retry
        from inventario.tasks import conciliar_conteos_nocturnos as task
        schedule=settings.CELERY_BEAT_SCHEDULE['inventario: conciliacion conteos nocturna']
        self.assertEqual(schedule['schedule'].hour,{2})
        self.assertEqual(settings.CELERY_TIMEZONE,'America/Mazatlan')
        with patch('inventario.conteos_nocturna.conciliar_conteos_nocturnos',return_value={'reintentar':True}), patch.object(task,'retry',side_effect=Retry()) as retry:
            with self.assertRaises(Retry):
                task.run()
            retry.assert_called_once_with(countdown=1800)

    def test_late_snapshot_retries_then_calculates_archived_history(self):
        self.snapshot.delete()
        self.acquire()
        first=conciliar_conteos_nocturnos()
        self.assertTrue(first['reintentar'])
        PointInventorySnapshot.objects.create(branch=self.point,product=self.product,stock=3,sync_job=self.job,
            captured_at=self.close+timedelta(minutes=30),raw_payload={'headers':['Codigo','Cantidad','Unidad'],'row':['C100','3','PZA']})
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            second=conciliar_conteos_nocturnos()
            client.assert_not_called()
        self.assertEqual(second['calculados'],1)

    def test_newer_cycle_can_resolve_saturated_coverage_without_changing_count(self):
        from contextlib import nullcontext
        boundary=movement(100,self.when-timedelta(hours=1),'VENTA','3','3')
        self.rows=[boundary]*497+self.rows
        self.fake.get_stock_history.return_value=self.rows
        self.acquire()
        self.job.started_at=self.when-timedelta(hours=8)
        self.job.finished_at=self.job.started_at+timedelta(minutes=2)
        self.job.save()
        self.snapshot.captured_at=self.job.started_at+timedelta(minutes=1)
        self.snapshot.save()
        with patch('inventario.conteos_nocturna.point_account_session_lock',return_value=nullcontext(False)):
            self.assertTrue(conciliar_conteos_nocturnos()['reintentar'])
        fresh=PointSyncJob.objects.create(status='SUCCESS',job_type='inventory',
            started_at=self.close+timedelta(minutes=29),finished_at=self.close+timedelta(minutes=31))
        PointInventorySnapshot.objects.create(branch=self.point,product=self.product,stock=3,sync_job=fresh,
            captured_at=self.close+timedelta(minutes=30),raw_payload=self.snapshot.raw_payload)
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            summary=conciliar_conteos_nocturnos()
            client.assert_not_called()
        self.assertEqual(summary['calculados'],1)
        self.assertEqual(self.line.lecturas.get().cantidad,3)

    def test_corrupt_archive_is_not_treated_as_complete(self):
        self.acquire()
        service=AuditStockHistoryService()
        record=service._existing_import(self.point,self.product)
        digest=record.raw_metadata['latest_response_fingerprint']
        record.raw_metadata['original_responses'][digest]['raw_sha256']='0'*64
        record.save(update_fields=['raw_metadata'])
        with self.assertRaises(AuditStockHistoryError):
            _historial(service,self.point,self.product)

    def test_duplicate_snapshot_in_one_cycle_is_pending_without_point_contact(self):
        self.acquire()
        self.snapshot.pk=None
        self.snapshot.save(force_insert=True)
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            summary=conciliar_conteos_nocturnos()
            client.assert_not_called()
        self.count.refresh_from_db()
        result=self.count.conciliacion['lineas'][str(self.line.pk)]
        self.assertIn('duplicado',result['motivo'])
        self.assertIsNone(result['diferencia'])
        self.assertTrue(summary['reintentar'])
        fresh=PointSyncJob.objects.create(status='SUCCESS',job_type='inventory',
            started_at=self.close+timedelta(minutes=39),finished_at=self.close+timedelta(minutes=41))
        PointInventorySnapshot.objects.create(branch=self.point,product=self.product,stock=3,sync_job=fresh,
            captured_at=self.close+timedelta(minutes=40),raw_payload=self.snapshot.raw_payload)
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            self.assertEqual(conciliar_conteos_nocturnos()['calculados'],1)
            client.assert_not_called()

    def test_incomplete_snapshot_retries_without_reading_point_until_unit_is_proven(self):
        self.acquire()
        self.snapshot.raw_payload['row'][2]=''
        self.snapshot.save()
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            self.assertTrue(conciliar_conteos_nocturnos()['reintentar'])
            client.assert_not_called()
        self.snapshot.raw_payload['row'][2]='PZA'
        self.snapshot.save()
        with patch('inventario.conteos_nocturna.PointHttpSessionClient') as client:
            self.assertEqual(conciliar_conteos_nocturnos()['calculados'],1)
            client.assert_not_called()

    def test_publish_rejects_recount_and_other_job_races(self):
        stale=ConteoSucursal.objects.get(pk=self.count.pk)
        readings=lecturas_efectivas(stale)
        result={'estado':'PENDIENTE','lineas':{},'huella_lecturas':huella_lecturas(readings)}
        self.assertTrue(_publicar(stale,readings,result))
        self.assertFalse(_publicar(stale,readings,{**result,'estado':'PARCIAL'}))
        self.count.refresh_from_db()
        self.action('reconteo',{'linea_ids':[self.line.pk],'motivo':'Verificar ubicación'},actor=self.admin)
        self.assertEqual(self.count.conciliacion,{})
        self.assertFalse(_publicar(stale,readings,result))
        self.assertEqual(self.line.lecturas.get(ronda=1).cantidad,3)

    def test_print_reuses_closing_comparison_without_mutations(self):
        self.acquire()
        conciliar_conteos_nocturnos()
        self.count.refresh_from_db()
        before = (self.count.version, self.count.conciliacion, self.count.eventos.count())
        self.client.force_login(self.admin)
        url = reverse('inventario:conteos_erp:detalle', args=[self.count.pk])
        self.assertContains(self.client.get(url), url + '?imprimir=1')
        response = self.client.get(url, {'imprimir': '1'})
        self.assertTemplateUsed(response, 'inventario/conteos/imprimir.html')
        self.assertContains(response, 'Físico al cierre')
        self.assertContains(response, '<td class="number">4</td>', count=2, html=True)
        self.assertContains(response, '<td class="number difference">0</td>', html=True)
        self.assertContains(response, 'Nota conservada')
        self.assertNotContains(response, 'Solicitar reconteo')
        self.assertIn('no-store', response['Cache-Control'])
        self.count.refresh_from_db()
        self.assertEqual(before, (self.count.version, self.count.conciliacion, self.count.eventos.count()))

    def test_print_pending_and_stale_results_never_show_zero_comparison(self):
        self.client.force_login(self.admin)
        url = reverse('inventario:conteos_erp:detalle', args=[self.count.pk])
        response = self.client.get(url, {'imprimir': '1'})
        self.assertContains(response, 'Pendiente de procesamiento nocturno.')
        self.assertContains(response, '<td class="number">—</td>', count=3, html=True)
        self.acquire()
        conciliar_conteos_nocturnos()
        self.count.refresh_from_db()
        self.count.conciliacion['huella_lecturas'] = 'obsolete'
        self.count.save(update_fields=['conciliacion'])
        response = self.client.get(url, {'imprimir': '1'})
        self.assertContains(response, '<td class="number">—</td>', count=3, html=True)
        self.assertNotContains(response, '<td class="number">4</td>', html=True)

    def test_print_does_not_expose_comparison_to_operator_or_other_branch(self):
        self.client.force_login(self.operator)
        url = reverse('operacion:conteos_app:detalle', args=[self.count.pk])
        self.assertNotContains(self.client.get(url), 'Imprimir / Guardar PDF')
        self.assertEqual(self.client.get(url, {'imprimir': '1'}).status_code, 403)
        outsider = get_user_model().objects.create_user('print-outsider')
        self.client.force_login(outsider)
        url = reverse('inventario:conteos_erp:detalle', args=[self.count.pk])
        self.assertEqual(self.client.get(url, {'imprimir': '1'}).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(url, {'imprimir': '1'}).status_code, 302)

    def test_comparison_is_hidden_in_recount_and_excel_preserves_system_history(self):
        self.acquire()
        conciliar_conteos_nocturnos()
        self.client.force_login(self.admin)
        detail=reverse('inventario:conteos_erp:detalle',args=[self.count.pk])
        self.assertContains(self.client.get(detail),'Físico al cierre')
        response=self.client.get(reverse('inventario:conteos_erp:exportar',args=[self.count.pk]))
        workbook=load_workbook(BytesIO(response.content))
        self.assertIn('Conciliación nocturna',workbook.sheetnames)
        self.assertIn('Sistema',str(list(workbook['Historial'].values)))
        self.action('reconteo',{'linea_ids':[self.line.pk],'motivo':'Revisión'},actor=self.admin)
        self.client.force_login(self.operator)
        response=self.client.get(reverse('operacion:conteos_app:detalle',args=[self.count.pk]))
        self.assertNotContains(response,'Físico al cierre')
        self.assertNotContains(response,'historial_sha256')

    def test_invalid_physical_time_does_not_block_quantity_or_export_as_formula(self):
        self.action('reconteo',{'linea_ids':[self.line.pk],'motivo':'Verificar'},actor=self.admin)
        self.action('iniciar',{})
        self.action('guardar',{'lecturas':{str(self.line.pk):{'cantidad':'1','contado_en':'=1+1'}}})
        reading=self.line.lecturas.get(ronda=2)
        self.assertEqual(reading.cantidad,1)
        self.assertIsNone(reading.contado_en)
        self.client.force_login(self.admin)
        response=self.client.get(reverse('inventario:conteos_erp:exportar',args=[self.count.pk]))
        workbook=load_workbook(BytesIO(response.content))
        cells=[row[9] for row in workbook['Lecturas históricas'].iter_rows(min_row=2)]
        preserved=next(cell for cell in cells if cell.value=="'=1+1")
        self.assertEqual(preserved.data_type,'s')
