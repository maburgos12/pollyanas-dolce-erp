from datetime import date
from uuid import uuid4
from tempfile import TemporaryDirectory
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from core.models import Sucursal, UserProfile
from pos_bridge.models import PointProduct


class ConteoRouteTests(SimpleTestCase):
    def test_app_and_erp_have_distinct_entrypoints(self):
        self.assertEqual(reverse('operacion:conteos_app:lista'), '/app/conteos/')
        self.assertEqual(reverse('inventario:conteos_erp:lista'), '/inventario/conteos-sucursales/')


class ConteoViewsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from inventario.services_conteos import preparar_conteo
        User = get_user_model()
        cls.admin = User.objects.create_superuser('coordinador', 'demo@example.invalid', 'test-local')
        cls.operator = User.objects.create_user('capturista', password='test-local')
        cls.other = User.objects.create_user('ajeno', password='test-local')
        cls.branch = Sucursal.objects.create(codigo='CT-M', nombre='Sucursal de prueba')
        UserProfile.objects.update_or_create(user=cls.operator, defaults={'sucursal':cls.branch})
        cls.product = PointProduct.objects.create(external_id='ct-100', sku='CT100', name='Producto conteo')
        cls.count = preparar_conteo(actor=cls.admin, sucursal=cls.branch, responsable=cls.operator, fecha=date.today(), titulo='Conteo cierre', items=[{'producto_id':cls.product.pk,'unidad':'PZA','fuente_unidad':'Reporte Point validado'}], request_id=str(uuid4()))

    def url(self, name='detalle'):
        return reverse('operacion:conteos_app:' + name, args=[self.count.pk])

    def test_assigned_capture_is_blind_no_reference_data(self):
        self.client.force_login(self.operator)
        self.count.referencia = {'secret_expected_stock':'314159', 'lineas':{}}
        self.count.save(update_fields=['referencia'])
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Conteo ciego')
        self.assertNotContains(response, '314159')
        self.assertIn('no-store', response['Cache-Control'])

    def test_other_branch_cannot_read_or_write(self):
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url()).status_code, 404)
        response = self.client.post(self.url('accion'), {'action':'cancelar','version':1,'request_id':str(uuid4()),'motivo':'intento ajeno'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 404)

    def test_mermas_only_user_can_capture_assigned_count_but_not_general_erp(self):
        from core.models import UserModuleAccess
        UserModuleAccess.objects.create(user=self.operator,module='mermas',access='manage')
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(self.url()).status_code,200)
        self.assertEqual(self.client.get('/inventario/').status_code,302)
        self.assertEqual(self.client.get('/inventario/')['Location'],'/mermas/app/')
        body={'action':'iniciar','version':1,'request_id':str(uuid4())}
        self.assertEqual(self.client.post(self.url('accion'),body,HTTP_X_REQUESTED_WITH='XMLHttpRequest').status_code,200)

    def test_csrf_required_for_mutations(self):
        from django.test import Client
        client=Client(enforce_csrf_checks=True)
        client.force_login(self.operator)
        response=client.post(self.url('accion'),{'action':'iniciar','version':1,'request_id':str(uuid4())})
        self.assertRedirects(response,'/login/',fetch_redirect_response=False)
        self.count.refresh_from_db()
        self.assertEqual(self.count.version,1)

    def test_zero_submit_and_replay(self):
        self.client.force_login(self.operator)
        begin=self.client.post(self.url('accion'), {'action':'iniciar','version':1,'request_id':str(uuid4())}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(begin.status_code,200,begin.content)
        self.assertIn('data-count-form',begin.json()['html'])
        line = self.count.lineas.get()
        body = {'action':'enviar','version':2,'request_id':str(uuid4()),f'cantidad_{line.pk}':'0','observaciones':'Cierre'}
        first = self.client.post(self.url('accion'), body, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(first.status_code, 200, first.content)
        again = self.client.post(self.url('accion'), body, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(again.status_code, 200, again.content)
        self.assertTrue(again.json()['ok'])
        self.count.refresh_from_db()
        self.assertEqual(self.count.estado, 'ENVIADO')

    def test_missing_cannot_submit(self):
        self.client.force_login(self.operator)
        self.client.post(self.url('accion'), {'action':'iniciar','version':1,'request_id':str(uuid4())}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        response = self.client.post(self.url('accion'), {'action':'enviar','version':2,'request_id':str(uuid4())}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Completa cada cantidad',response.json()['toast']['message'])

    def test_capture_operator_cannot_export_expected_values(self):
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(self.url('exportar')).status_code, 403)

    def test_reviewer_can_open_submitted_count_without_reference_and_export(self):
        from inventario.services_conteos import ejecutar_accion
        from openpyxl import load_workbook
        from io import BytesIO
        for version,action,payload in [(1,'iniciar',{}),(2,'enviar',{'lecturas':{str(self.count.lineas.get().pk):{'cantidad':'0','incidencia':''}}})]:
            ejecutar_accion(conteo_id=self.count.pk,actor=self.operator,action=action,version=version,request_id=uuid4(),payload=payload)
        self.client.force_login(self.admin)
        response=self.client.get(self.url())
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'Sin referencia')
        self.assertContains(response,'Solicitar reconteo')
        exported=self.client.get(self.url('exportar'))
        self.assertEqual(exported.status_code,200)
        row=list(load_workbook(BytesIO(exported.content)).active.values)[1]
        self.assertEqual(row[8],0)
        self.assertIsNone(row[10])
        self.assertIsNone(row[11])
        book=load_workbook(BytesIO(exported.content))
        self.assertIn('Lecturas históricas',book.sheetnames)
        self.assertIn('Referencias históricas',book.sheetnames)
        self.assertEqual(book['Lecturas históricas'].max_row,2)
        self.assertEqual(book['Lecturas históricas'].cell(2,8).value,'0.000000')

    def test_get_never_changes_version(self):
        self.client.force_login(self.admin)
        version = self.count.version
        self.client.get(reverse('inventario:conteos_erp:detalle', args=[self.count.pk]))
        self.count.refresh_from_db()
        self.assertEqual(self.count.version, version)

    def test_large_scope_uses_compact_payload_below_django_field_limit(self):
        import json
        from inventario.models_conteos import ConteoSucursal
        products=PointProduct.objects.bulk_create([PointProduct(external_id=f'bulk-{i}',sku=f'B{i}',name=f'Artículo {i}') for i in range(510)])
        items=[{'producto_id':p.pk,'unidad':'PZA','fuente_unidad':'Prueba de formato compacto'} for p in products]
        self.client.force_login(self.admin)
        response=self.client.post(reverse('operacion:conteos_app:preparar'),{'sucursal':self.branch.pk,'responsable':self.operator.pk,'fecha':date.today().isoformat(),'titulo':'Alcance grande','request_id':str(uuid4()),'articulos_json':json.dumps(items)},HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,200,response.content)
        count=ConteoSucursal.objects.get(titulo='Alcance grande')
        url=reverse('operacion:conteos_app:accion',args=[count.pk])
        self.client.force_login(self.operator)
        self.assertEqual(self.client.post(url,{'action':'iniciar','version':1,'request_id':str(uuid4())},HTTP_X_REQUESTED_WITH='XMLHttpRequest').status_code,200)
        response=self.client.post(url,{'action':'enviar','version':2,'request_id':str(uuid4()),'lecturas_json':json.dumps({str(line.pk):{'cantidad':'0','incidencia':''} for line in count.lineas.all()})},HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,200,response.content)
        self.assertEqual(count.lineas.filter(lecturas__cantidad=0).count(),510)

    def test_evidence_download_requires_count_access(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        directory=TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        settings=override_settings(CONTEOS_PRIVATE_ROOT=directory.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.client.force_login(self.operator)
        file=SimpleUploadedFile('observacion.pdf',b'%PDF-1.4\nExample evidence',content_type='application/pdf')
        response=self.client.post(self.url('evidencia'),{'version':1,'request_id':str(uuid4()),'archivo':file},HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code,200,response.content)
        self.assertEqual(response.json()['target'],'#conteo-detail')
        self.assertIn('observacion.pdf',response.json()['html'])
        self.assertIn('data-version="2"',response.json()['html'])
        event=self.count.eventos.get(action='evidencia')
        url=reverse('operacion:conteos_app:descarga',args=[self.count.pk,event.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(url).status_code,404)
