"""App y ERP comparten el servicio; ninguna vista escribe inventario operativo."""
from decimal import Decimal
from copy import copy
import json
from io import BytesIO
from uuid import uuid4

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook

from maestros.models import Insumo
from pos_bridge.models import PointProduct
from .conteos_access import conteos_visibles, puede_capturar, puede_coordinar, puede_revisar
from .forms_conteos import PrepararConteoForm
from .models_conteos import LecturaConteoSucursal
from .services_conteos import ConteoConflict, ejecutar_accion, lecturas_efectivas, preparar_conteo


def _namespace(request):
    return 'operacion:conteos_app' if request.path.startswith('/app/') else 'inventario:conteos_erp'


def _base(request):
    return 'inventario/conteos/base_app.html' if _namespace(request) == 'operacion:conteos_app' else 'base.html'


def _url(request, name, *args):
    return reverse(_namespace(request) + ':' + name, args=args)


def _count(request, pk):
    return get_object_or_404(conteos_visibles(request.user).select_related('sucursal','responsable'), pk=pk)


def _context(request, count, *, ack='', ack_version=None):
    reviewer = puede_revisar(request.user, count)
    editable = puede_capturar(request.user, count) and count.estado in ('CAPTURA','RECONTEO')
    readings = (list(LecturaConteoSucursal.objects.filter(linea__conteo=count, ronda=count.ronda).select_related('linea').order_by('linea__nombre')) if editable else lecturas_efectivas(count))
    rows = []
    reference = count.referencia if reviewer and not editable else {}
    for reading in readings:
        line = reading.linea
        source = reference.get('lineas', {}).get(str(line.pk), {})
        difference = None
        if source.get('cantidad') is not None and reading.cantidad is not None:
            difference = reading.cantidad - Decimal(source['cantidad'])
        rows.append({'lectura':reading, 'linea':line, 'referencia':source, 'diferencia':difference})
    captured = sum(r.cantidad is not None or bool(r.incidencia) for r in readings)
    # Event payloads can contain expected quantities and previous rounds. Never
    # serialize those into a blind capture page (even hidden DOM/JSON).
    events = count.eventos.select_related('actor').order_by('-pk') if reviewer and not editable else []
    history = []
    for event in events:
        payload = event.payload
        summary = payload.get('motivo', '')
        snapshot = payload.get('snapshot', [])
        historical_rows = []
        names = {str(row['linea'].pk):row['linea'].nombre for row in rows}
        for item in snapshot:
            historical_rows.append({'nombre':names.get(str(item['linea_id']), str(item['linea_id'])), **item})
        history.append({'evento':event,'motivo':summary,'lecturas':historical_rows})
    return {'base_template':_base(request),'conteo':count,'rows':rows,'editable':editable,
            'revisor':reviewer,'referencia':reference,'historial':history,
            'motivo_reconteo':(count.eventos.filter(action='reconteo').order_by('-pk').values_list('payload__motivo',flat=True).first() or '') if editable and count.ronda > 1 else '',
            'capturados':captured,'total':len(readings),'avance':round(captured*100/len(readings)) if readings else 0,
            'request_id':str(uuid4()),'action_ids':{key:str(uuid4()) for key in ('start','reference','recount','accept','verify','evidence','cancel')},'ack':ack,'ack_version':ack_version,
            'lista_url':_url(request,'lista'),'accion_url':_url(request,'accion',count.pk),
            'detalle_url':_url(request,'detalle',count.pk),'export_url':_url(request,'exportar',count.pk),
            'evidencia_url':_url(request,'evidencia',count.pk), 'usuario_id':request.user.pk,
            'evidencias':_evidence_links(request,count,editable)}


def _evidence_links(request,count,editable):
    qs = count.eventos.filter(action='evidencia')
    if editable:
        qs = qs.filter(payload__ronda=count.ronda)
    return [{'nombre':e.payload.get('nombre','Evidencia'),'url':_url(request,'descarga',count.pk,e.pk)} for e in qs]


@login_required
@never_cache
@require_GET
def lista(request):
    qs = conteos_visibles(request.user).select_related('sucursal','responsable')
    branch = request.GET.get('sucursal','')
    state = request.GET.get('estado','')
    if branch.isdigit(): qs = qs.filter(sucursal_id=branch)
    if state in ('CAPTURA','RECONTEO','ENVIADO','ACEPTADO','CANCELADO'): qs = qs.filter(estado=state)
    from django.core.paginator import Paginator
    page = Paginator(qs,30).get_page(request.GET.get('page'))
    return render(request,'inventario/conteos/lista.html',{'base_template':_base(request),'page':page,
        'rows':[{'conteo':c,'url':_url(request,'detalle',c.pk)} for c in page],
        'puede_preparar':puede_coordinar(request.user),'preparar_url':_url(request,'preparar'),
        'estado':state,'sucursal':branch})


@login_required
@never_cache
@require_GET
def detalle(request, pk):
    return render(request,'inventario/conteos/detalle.html',_context(request,_count(request,pk)))


def _payload(request, action):
    if action in ('guardar','enviar'):
        if 'lecturas_json' in request.POST:
            return {'lecturas':json.loads(request.POST['lecturas_json']),'observaciones':request.POST.get('observaciones','')}
        ids = {key.removeprefix('cantidad_') for key in request.POST if key.startswith('cantidad_')}
        ids.update(key.removeprefix('incidencia_') for key in request.POST if key.startswith('incidencia_'))
        return {'lecturas':{key:{'cantidad':request.POST.get('cantidad_'+key,''),'incidencia':request.POST.get('incidencia_'+key,'')} for key in ids},'observaciones':request.POST.get('observaciones','')}
    if action == 'reconteo':
        return {'linea_ids':[int(i) for i in request.POST.getlist('linea_ids')], 'motivo':request.POST.get('motivo','')}
    if action in ('referencia','iniciar'): return {}
    if action == 'validar_referencia': return {'motivo':request.POST.get('motivo',''), 'confirmado':request.POST.get('confirmado') == '1'}
    return {'motivo':request.POST.get('motivo','')}


def _async(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest' or 'application/json' in request.headers.get('Accept','')


def _error(request, count, error, status=400):
    text = ' '.join(error.messages) if isinstance(error,ValidationError) else str(error)
    if _async(request): return JsonResponse({'ok':False,'toast':{'type':'error','message':text,'persistent':True}},status=status)
    context = _context(request,count)
    if context['editable'] and request.POST.get('action') in ('guardar','enviar'):
        for row in context['rows']:
            row['posted_cantidad'] = request.POST.get(f"cantidad_{row['linea'].pk}", '')
            row['posted_incidencia'] = request.POST.get(f"incidencia_{row['linea'].pk}", '')
        context['posted'] = True
        context['posted_observaciones'] = request.POST.get('observaciones','')
    messages.error(request,text)
    return render(request,'inventario/conteos/detalle.html',context,status=status)


@login_required
@never_cache
@require_POST
def accion(request, pk):
    count = _count(request,pk)
    action = request.POST.get('action','')
    authorized = puede_capturar(request.user,count) if action in ('iniciar','guardar','enviar') else puede_revisar(request.user,count)
    if not authorized: return _error(request,count,'No tienes permiso para esta acción.',403)
    try:
        old_version = int(request.POST.get('version',''))
        result = ejecutar_accion(conteo_id=count.pk, actor=request.user, action=action, version=old_version,
            request_id=request.POST.get('request_id',''), payload=_payload(request,action))
    except ConteoConflict as exc: return _error(request,count,exc,409)
    except (ValidationError,ValueError,TypeError) as exc: return _error(request,count,exc)
    count.refresh_from_db()
    notice = {'iniciar':'Inicio del conteo registrado. Ya puedes capturar.','guardar':'Avance guardado.','enviar':'Conteo enviado a revisión.','reconteo':'Reconteo solicitado. Se conserva la captura anterior.',
              'aceptar':'Conteo físico aceptado. No se modificaron existencias.','cancelar':'Conteo cancelado. Se conserva su historial.',
              'referencia':'Referencia Point registrada. Revisa fecha, unidad y cobertura.',
              'validar_referencia':'Declaración de revisión del corte registrada.'}.get(action,'Acción registrada.')
    if _async(request):
        html = render_to_string('inventario/conteos/_detalle.html',_context(request,count,ack=request.POST['request_id'],ack_version=old_version),request=request)
        return JsonResponse({'ok':True,'result':result,'target':'#conteo-detail','html':html,'toast':{'type':'success','message':notice}})
    messages.success(request,notice)
    return redirect(_url(request,'detalle',count.pk)+'#conteo-detail')


def _catalogo(tipo, query):
    if tipo == 'insumo':
        qs = Insumo.objects.filter(activo=True).exclude(codigo_point='').select_related('unidad_base')
        if query: qs=qs.filter(Q(nombre__icontains=query)|Q(codigo_point__icontains=query))
        return [{'key':f'i{x.pk}','nombre':x.nombre,'codigo':x.codigo_point,'unidad':x.unidad_base.codigo if x.unidad_base_id else '',
                 'fuente':'Unidad base del catálogo canónico' if x.unidad_base_id else ''} for x in qs.order_by('nombre')[:200]]
    qs = PointProduct.objects.filter(active=True).exclude(sku='').exclude(sku__in=Insumo.objects.filter(activo=True).exclude(codigo_point='').values('codigo_point'))
    if query: qs=qs.filter(Q(name__icontains=query)|Q(sku__icontains=query)|Q(category__icontains=query))
    return [{'key':f'p{x.pk}','nombre':x.name,'codigo':x.sku,'unidad':'','fuente':''} for x in qs.order_by('name')[:200]]


@login_required
@never_cache
def preparar(request):
    if not puede_coordinar(request.user): raise PermissionDenied
    if request.method not in ('GET','POST'): return HttpResponse(status=405)
    form = PrepararConteoForm(request.POST or None, initial={'fecha':timezone.localdate(),'request_id':uuid4()})
    tipo = request.GET.get('tipo','producto')
    query = request.GET.get('q','').strip()[:120]
    if request.method == 'POST' and form.is_valid():
        try:
            if 'articulos_json' in request.POST:
                items=json.loads(request.POST['articulos_json'])
            else:
                items = []
                for key in request.POST.getlist('articulos'):
                    if len(key)<2 or key[0] not in 'pi' or not key[1:].isdigit(): raise ValidationError('Artículo inválido.')
                    items.append({'producto_id' if key[0]=='p' else 'insumo_id':int(key[1:]),
                                  'unidad':request.POST.get('unidad_'+key,''), 'fuente_unidad':request.POST.get('fuente_'+key,'')})
            count = preparar_conteo(actor=request.user,items=items,**form.cleaned_data)
        except (ValidationError,ValueError) as exc:
            form.add_error(None,exc)
        else:
            url = _url(request,'detalle',count.pk)
            if _async(request): return JsonResponse({'ok':True,'redirect':url,'toast':{'type':'success','message':'Conteo preparado y asignado.'}})
            messages.success(request,'Conteo preparado y asignado.')
            return redirect(url)
    catalog = _catalogo(tipo,query)
    if request.method == 'POST':
        for row in catalog:
            row['selected'] = row['key'] in request.POST.getlist('articulos')
            row['unidad'] = request.POST.get('unidad_'+row['key'],row['unidad'])
            row['fuente'] = request.POST.get('fuente_'+row['key'],row['fuente'])
        if _async(request):
            return JsonResponse({'ok':False,'toast':{'type':'error','message':' '.join(str(e) for errors in form.errors.values() for e in errors),'persistent':True}},status=400)
    return render(request,'inventario/conteos/preparar.html',{'base_template':_base(request),'form':form,'catalogo':catalog,'tipo':tipo,'q':query,'lista_url':_url(request,'lista')})


def _excel_text(value):
    value = '' if value is None else str(value)
    return "'"+value if value.lstrip().startswith(('=','+','-','@')) else value


def _excel_number(value):
    if value is None: return None
    return str(value) if len(value.normalize().as_tuple().digits)>15 else float(value)


@login_required
@never_cache
@require_GET
def exportar(request, pk):
    count = _count(request,pk)
    if not puede_revisar(request.user,count): raise PermissionDenied
    wb = Workbook()
    sheet=wb.active
    sheet.title='Conteo físico'
    sheet.append(['Folio','Sucursal','Fecha programada','Estado','Código','Artículo','Unidad','Ronda','Contado','Incidencia','Referencia Point','Diferencia informativa','Extracción','Conciliación','Inicio ronda actual','Último envío','Observaciones','Fuente de unidad','Ciclo Point','Snapshot Point','Lectura actualizada'])
    for r in lecturas_efectivas(count):
        line=r.linea
        ref=count.referencia.get('lineas',{}).get(str(line.pk),{})
        qty=Decimal(ref['cantidad']) if ref.get('cantidad') is not None else None
        sheet.append([f'CF-{count.pk}',_excel_text(count.sucursal.nombre),count.fecha.isoformat(),count.estado,_excel_text(line.codigo),_excel_text(line.nombre),_excel_text(line.unidad),r.ronda,
                      _excel_number(r.cantidad),_excel_text(r.incidencia),_excel_number(qty),
                      _excel_number(r.cantidad-qty) if r.cantidad is not None and qty is not None else None,ref.get('capturado_en',''),
                      'Declaración de corte por revisor' if count.referencia.get('corte_verificado') else 'Pendiente de conciliación temporal',
                      count.iniciado_en.isoformat() if count.iniciado_en else '',count.enviado_en.isoformat() if count.enviado_en else '',
                      _excel_text(count.observaciones),_excel_text(line.fuente_unidad),ref.get('sync_job_id'),ref.get('snapshot_id'),r.actualizado_en.isoformat()])
    history=wb.create_sheet('Historial')
    history.append(['Fecha','Acción','Usuario','Detalle'])
    readings=wb.create_sheet('Lecturas históricas')
    readings.append(['Evento','Fecha','Acción','Usuario','Artículo','Código','Ronda','Cantidad','Incidencia'])
    references=wb.create_sheet('Referencias históricas')
    references.append(['Evento','Fecha','Artículo','Código','Cantidad','Unidad','Ciclo Point','Snapshot Point','Extracción','Estado','Corte declarado','Error'])
    lines={str(line.pk):line for line in count.lineas.all()}
    import json
    for e in count.eventos.select_related('actor'):
        safe_payload={k:v for k,v in e.payload.items() if k not in {'archivo','snapshot','lecturas','referencia','articulos'}}
        reference=e.payload.get('referencia',{})
        if reference:
            safe_payload['referencia']={k:v for k,v in reference.items() if k!='lineas'}
        historical=e.payload.get('snapshot') or [{'linea_id':key,'ronda':e.payload.get('resultado',{}).get('ronda'),**value} for key,value in e.payload.get('lecturas',{}).items()]
        for item in historical:
            line=lines.get(str(item['linea_id']))
            readings.append([e.pk,e.creado_en.isoformat(),e.action,_excel_text(e.actor.username),_excel_text(line.nombre if line else ''),_excel_text(line.codigo if line else ''),item.get('ronda'),_excel_text(item.get('cantidad')),_excel_text(item.get('incidencia'))])
        for key,item in reference.get('lineas',{}).items():
            line=lines.get(str(key))
            references.append([e.pk,e.creado_en.isoformat(),_excel_text(line.nombre if line else ''),_excel_text(line.codigo if line else ''),_excel_text(item.get('cantidad')),_excel_text(item.get('unidad')),item.get('sync_job_id'),item.get('snapshot_id'),item.get('capturado_en'),reference.get('estado'),bool(reference.get('corte_verificado')),_excel_text(item.get('error'))])
        history.append([e.creado_en.isoformat(),e.action,_excel_text(e.actor.username),_excel_text(json.dumps(safe_payload,ensure_ascii=False))])
    for ws in wb:
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width=min(48,max(14,len(str(column[0].value))+2))
    output=BytesIO(); wb.save(output)
    response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="conteo-CF-{count.pk}.xlsx"'
    return response


@login_required
@never_cache
@require_POST
def evidencia(request, pk):
    from .conteos_evidence import adjuntar_evidencia
    count = _count(request,pk)
    if not (puede_revisar(request.user,count) or puede_capturar(request.user,count)):
        return _error(request,count,'No tienes permiso para adjuntar evidencia.',403)
    try:
        result=adjuntar_evidencia(conteo_id=count.pk,actor=request.user,version=int(request.POST.get('version','')),request_id=request.POST.get('request_id',''),uploaded=request.FILES.get('archivo'))
    except ConteoConflict as exc: return _error(request,count,exc,409)
    except (ValidationError,ValueError,TypeError) as exc: return _error(request,count,exc)
    if _async(request):
        count.refresh_from_db()
        html = render_to_string('inventario/conteos/_detalle.html',_context(request,count),request=request)
        return JsonResponse({'ok':True,'result':result,'target':'#conteo-detail','html':html,'toast':{'type':'success','message':'Evidencia recibida.'}})
    return redirect(_url(request,'detalle',count.pk)+'#evidencias')


@login_required
@never_cache
@require_GET
def descarga(request, pk, event_id):
    from .conteos_evidence import descargar_evidencia
    return descargar_evidencia(_count(request,pk),event_id,request.user)
