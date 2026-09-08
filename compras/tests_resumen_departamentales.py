from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, SimpleTestCase, RequestFactory
from django.urls import reverse
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import load_workbook

from maestros.models import Proveedor
from reportes.models import AreaPresupuesto
from .models import (
    SolicitudCompraDepartamental as Solicitud,
    ItemCompraDepartamental as Item,
    CotizacionCompraDepartamental as Cotizacion,
    CompromisoCompraDepartamental as Compromiso,
)


class ResumenDepartamentalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser('resumen', password='test')
        cls.area = AreaPresupuesto.objects.create(codigo='resumen-ventas', nombre='Ventas resumen')
        cls.otra = AreaPresupuesto.objects.create(codigo='resumen-produccion', nombre='Producción resumen')
        cls.proveedor = Proveedor.objects.create(nombre='Proveedor resumen')

    def setUp(self):
        self.client.force_login(self.user)
        self.url = reverse('compras:departamental_bandeja')

    def solicitud(self, **kwargs):
        return Solicitud.objects.create(**{
            'area': self.area, 'solicitante': self.user, 'periodo': date(2026, 9, 1),
            'estado': Solicitud.ESTADO_ENVIADA, **kwargs,
        })

    def item(self, solicitud=None, **kwargs):
        return Item.objects.create(**{
            'solicitud': solicitud or self.solicitud(), 'descripcion': 'Equipo',
            'cantidad': Decimal('2'), 'costo_unitario_estimado': Decimal('100.25'), **kwargs,
        })

    def cotizacion(self, item, **kwargs):
        return Cotizacion.objects.create(**{
            'item': item, 'proveedor': self.proveedor, 'cantidad_ofertada': Decimal('2'),
            'costo_unitario': Decimal('110'), 'seleccionada': True,
            'impuestos': Decimal('35.20'), 'envio': Decimal('20'), 'descuento': Decimal('5'), **kwargs,
        })

    def test_excluye_borradores_canceladas_completadas_y_articulos_terminales(self):
        visible = self.item()
        for estado in (Solicitud.ESTADO_BORRADOR, Solicitud.ESTADO_CANCELADA, Solicitud.ESTADO_COMPLETADA):
            self.item(self.solicitud(estado=estado))
        for estado in (Item.ESTADO_CANCELADO, Item.ESTADO_RECHAZADO, Item.ESTADO_RECIBIDO_CONFORME):
            self.item(estado=estado)
        response = self.client.get(self.url)
        self.assertEqual([i.pk for i in response.context['items']], [visible.pk])
        self.assertEqual(response.context['resumen']['solicitado'], Decimal('200.50'))

    def test_totales_sin_duplicar_y_cobertura_de_precios(self):
        solicitud = self.solicitud()
        item = self.item(solicitud)
        quote = self.cotizacion(item)
        self.cotizacion(item, seleccionada=False, costo_unitario=999)
        Compromiso.objects.create(item=item, cotizacion=quote, monto=Decimal('270.20'), formalizado_en=timezone.now())
        sin_estimacion = self.item(solicitud, costo_unitario_estimado=None)
        quote2 = self.cotizacion(sin_estimacion)
        Compromiso.objects.create(item=sin_estimacion, cotizacion=quote2, monto=999)  # reserva, no orden
        self.item(solicitud, costo_unitario_estimado=None)
        self.item(solicitud, costo_unitario_estimado=0)
        response = self.client.get(self.url)
        total = response.context['resumen']
        self.assertEqual(total['solicitudes'], 1)
        self.assertEqual(total['articulos'], 4)
        self.assertEqual(total['solicitado'], Decimal('200.50'))
        self.assertEqual(total['cotizado'], Decimal('540.40'))
        self.assertEqual(total['comprometido'], Decimal('270.20'))
        self.assertEqual(total['sin_estimacion'], 2)
        self.assertEqual(total['sin_precio'], 1)
        self.assertEqual(total['sin_cotizacion'], 2)
        self.assertContains(response, 'Total estimado parcial')

    def test_filtros_combinados_y_enlace_por_departamento(self):
        visible = self.item(estado=Item.ESTADO_POR_COTIZAR)
        self.item(self.solicitud(area=self.otra), estado=Item.ESTADO_POR_COTIZAR)
        self.item(self.solicitud(periodo=date(2026, 8, 1)), estado=Item.ESTADO_POR_COTIZAR)
        self.item()
        params = {'area': self.area.pk, 'periodo': '2026-09', 'estado': Item.ESTADO_POR_COTIZAR}
        response = self.client.get(self.url, params)
        self.assertEqual([i.pk for i in response.context['items']], [visible.pk])
        self.assertEqual(response.context['departamentos'][0]['solicitado'], Decimal('200.50'))
        self.assertIn('periodo=2026-09', response.context['departamentos'][0]['url'])

    def test_filtros_invalidos_no_amplian_consulta(self):
        self.item()
        for params in ({'periodo': '2026-99'}, {'area': '999999999'}, {'estado': 'CANCELADO'}):
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertTrue(response.context['filtros'].errors)
                self.assertEqual(response.context['resumen']['articulos'], 0)
                response = self.client.get(self.url, {**params, 'exportar': 'xlsx'})
                self.assertEqual(response.status_code, 400)

    def test_excel_mismo_filtro_importes_numericos_y_texto_sin_formulas(self):
        self.item(descripcion='=1+1')
        self.item(self.solicitud(area=self.otra))
        response = self.client.get(self.url, {'area': self.area.pk, 'exportar': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Resumen', 'Artículos'])
        rows = list(wb['Artículos'].values)
        header = next(row for row in rows if row[0] == 'Solicitud')
        data = rows[rows.index(header) + 1:]
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0][header.index('Solicitado estimado')], 200.50)
        self.assertEqual(data[0][header.index('Cotizado')], None)
        self.assertFalse(any(cell.data_type == 'f' for ws in wb for row in ws for cell in row))

    def test_compromiso_inactivo_no_se_suma_y_centavos_concilian(self):
        item = self.item(cantidad=Decimal('0.333'), costo_unitario_estimado=Decimal('1.01'))
        quote = self.cotizacion(item)
        Compromiso.objects.create(item=item, cotizacion=quote, monto=100, activo=False, formalizado_en=timezone.now())
        response = self.client.get(self.url)
        self.assertEqual(response.context['resumen']['solicitado'], Decimal('0.34'))
        self.assertEqual(response.context['resumen']['comprometido'], 0)

    def test_acceso_existente_tambien_protege_exportacion(self):
        user = get_user_model().objects.create_user('sin-acceso')
        self.client.force_login(user)
        for params in ({}, {'exportar': 'xlsx'}):
            self.assertEqual(self.client.get(self.url, params).status_code, 403)

    def test_consulta_no_crece_por_articulo(self):
        from .resumen_departamentales import construir_resumen_departamental
        for _ in range(5):
            self.cotizacion(self.item())
        with self.assertNumQueries(2):
            contexto = construir_resumen_departamental({})
        self.assertEqual(contexto['resumen']['articulos'], 5)

    def test_vacio_es_cero_sin_advertencia_de_datos_faltantes(self):
        response = self.client.get(self.url)
        self.assertEqual(response.context['resumen']['articulos'], 0)
        self.assertEqual(response.context['resumen']['solicitado'], 0)
        self.assertNotContains(response, 'Total estimado parcial')



class ResumenTemplateTests(SimpleTestCase):
    def test_bandeja_expone_resumen_y_filtros(self):
        from django.contrib.auth.models import AnonymousUser
        request = RequestFactory().get('/')
        request.user = AnonymousUser()
        html = render_to_string('compras/departamentales/bandeja.html', {'items': [], 'request': request})
        self.assertIn('Solicitado estimado', html)
        self.assertIn('Exportar a Excel', html)
        self.assertIn('Por departamento', html)
