"""Regresiones de identidad e historial entre compras sucesivas del mismo artículo."""
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from compras.models import (CotizacionCompraDepartamental,
                            RecepcionItemDepartamental, ReembolsoCompraDepartamental)
from compras.services_departamentales import (confirmar_recepcion_departamental,
    evaluar_presupuesto_item, generar_ordenes_departamentales, seleccionar_cotizacion)
from compras.services_intentos_compra import cancelar_intento_compra, registrar_reembolso_compra
from compras.services_edicion_compra import editar_cotizacion
from compras.tests_edicion_compra import _CompraDepartamentalBase
from compras.resumen_departamentales import construir_resumen_departamental
from maestros.models import Proveedor


class CrucesIntentosCompraTests(_CompraDepartamentalBase, TestCase):
    def cancelar(self, intento, *, pagado=False):
        return cancelar_intento_compra(
            intento, version=intento.version, motivo='NO_ENTREGO', detalle='No entregará', actor=self.user,
            **({'reembolso_solicitado': Decimal('1000'),
                'reembolso_solicitado_en': timezone.localdate()} if pagado else {}),
        )

    def reemplazar(self):
        response = self.client.post(reverse('compras:departamental_cotizar', args=[self.item.pk]), {
            'proveedor': self.proveedor.pk, 'cantidad_ofertada': '2', 'costo_unitario': '100',
            'descuento': '0', 'impuestos': '0', 'envio': '0', 'instalacion': '0', 'otros_cargos': '0',
        }, **self.headers)
        self.assertEqual(response.status_code, 200, response.content)
        quote = self.item.cotizaciones.order_by('-pk').first()
        seleccionar_cotizacion(quote, actor=self.user)
        generar_ordenes_departamentales([self.item], actor=self.user)
        return self.item.intento_vigente

    def recibir(self, intento, cantidad='1', **cambios):
        return self.client.post(reverse('compras:departamental_recibir', args=[self.item.pk]), {
            'cantidad_recibida': cantidad, 'intento_id': intento.pk,
            'intento_version': intento.version, **cambios,
        }, **self.headers)

    def test_cotizacion_cancelada_sin_pago_es_inmutable_y_admite_nueva(self):
        self.ordenar()
        intento = self.item.intento_vigente
        self.cancelar(intento)
        self.verificar_cotizacion_historica(intento)
        self.assertNotEqual(self.reemplazar().cotizacion_id, self.quote.pk)

    def test_cotizacion_reembolso_es_inmutable_y_admite_nueva(self):
        self.quote.costo_unitario = 500
        self.quote.save(update_fields=['costo_unitario'])
        self.assertEqual(self.comprar(importe_final='1000').status_code, 200)
        intento = self.item.intento_vigente
        self.cancelar(intento, pagado=True)
        self.verificar_cotizacion_historica(intento)
        self.assertNotEqual(self.reemplazar().cotizacion_id, self.quote.pk)

    def verificar_cotizacion_historica(self, intento):
        self.quote.refresh_from_db()
        antes = (self.quote.proveedor_id, self.quote.cantidad_ofertada, self.quote.costo_unitario,
                 self.quote.documento.name, self.quote.version)
        linea = intento.linea_orden
        proveedor = Proveedor.objects.create(nombre='Proveedor nuevo')
        response = self.editar(proveedor=proveedor.pk, cantidad_ofertada='4', costo_unitario='999')
        self.assertEqual(response.status_code, 409, response.content)
        self.assertEqual(self.editar(documento=SimpleUploadedFile(
            'reemplazo.pdf', b'%PDF-1.4\n%%EOF', content_type='application/pdf',
        )).status_code, 409)
        with self.assertRaises(ValidationError):
            editar_cotizacion(self.quote, datos={'costo_unitario': Decimal('999')},
                              version=self.quote.version, motivo='Cambio directo', actor=self.user)
        self.quote.refresh_from_db(); linea.refresh_from_db(); intento.refresh_from_db()
        self.assertEqual(antes, (self.quote.proveedor_id, self.quote.cantidad_ofertada,
                                self.quote.costo_unitario, self.quote.documento.name, self.quote.version))
        self.assertEqual(linea.orden.proveedor_id, antes[0])
        self.assertEqual(intento.cotizacion_id, self.quote.pk)
        detalle = self.client.get(reverse('compras:departamental_detalle', args=[self.solicitud.pk]))
        self.assertNotContains(detalle, reverse('compras:departamental_cotizacion_editar', args=[self.quote.pk]))

    def test_recepcion_pestana_antigua_no_se_aplica_al_reemplazo(self):
        self.ordenar()
        anterior = self.item.intento_vigente
        self.cancelar(anterior)
        nuevo = self.reemplazar()
        response = self.recibir(anterior)
        self.assertEqual(response.status_code, 409, response.content)
        self.assertFalse(RecepcionItemDepartamental.objects.filter(linea_orden=nuevo.linea_orden).exists())
        self.assertIn(f'#item-{self.item.pk}', response.json()['redirect'])

    def test_recepcion_requiere_identidad_y_version(self):
        self.ordenar()
        intento = self.item.intento_vigente
        for datos in ({'intento_id': ''}, {'intento_version': ''}, {'intento_version': 0}):
            with self.subTest(datos=datos):
                self.assertEqual(self.recibir(intento, **datos).status_code, 409)
        self.assertFalse(RecepcionItemDepartamental.objects.exists())

    def test_recepcion_parcial_invalida_doble_envio_y_form_actualiza_version(self):
        self.ordenar()
        intento = self.item.intento_vigente
        self.assertEqual(self.recibir(intento).status_code, 200)
        self.assertEqual(self.recibir(intento).status_code, 409)
        intento.refresh_from_db()
        self.assertEqual(intento.version, 2)
        detalle = self.client.get(reverse('compras:departamental_detalle', args=[self.solicitud.pk]))
        self.assertContains(detalle, f'name="intento_id" value="{intento.pk}"')
        self.assertContains(detalle, 'name="intento_version" value="2"')
        self.assertEqual(self.recibir(intento).status_code, 200)
        self.assertEqual(RecepcionItemDepartamental.objects.count(), 2)

    def test_correccion_cantidad_invalida_version_sin_impedir_reabrir_parcial(self):
        self.item.cantidad = 4
        self.item.save(update_fields=['cantidad'])
        self.ordenar()
        intento = self.item.intento_vigente
        recepcion = RecepcionItemDepartamental.objects.create(
            linea_orden=intento.linea_orden, cantidad_recibida=4, registrado_por=self.user)
        intento.refresh_from_db()
        self.assertEqual((intento.estado, intento.version), ('ENTREGADO', 2))
        recepcion.cantidad_recibida = 3
        recepcion.save(update_fields=['cantidad_recibida'])
        intento.refresh_from_db()
        self.assertEqual((intento.estado, intento.version), ('VIGENTE', 3))
        recepcion.observaciones = 'Corrección documentada'
        recepcion.save(update_fields=['observaciones'])
        intento.refresh_from_db()
        self.assertEqual(intento.version, 3)

    def test_detalle_entregado_no_cuenta_dos_veces_y_otro_item_conserva_exposicion(self):
        self.ordenar()
        intento = self.item.intento_vigente
        RecepcionItemDepartamental.objects.create(
            linea_orden=intento.linea_orden, cantidad_recibida=2, registrado_por=self.user)
        detalle = self.client.get(reverse('compras:departamental_detalle', args=[self.solicitud.pk]))
        visible = list(detalle.context['solicitud'].items.all())[0]
        self.assertEqual(visible.evaluacion_presupuesto.disponible_antes, Decimal('10000'))
        self.assertEqual(visible.evaluacion_presupuesto.disponible_despues, Decimal('9800'))
        otro = self.solicitud.items.create(descripcion='Otro', cantidad=1, rubro=self.rubro)
        evaluacion = evaluar_presupuesto_item(otro, Decimal('100'))
        self.assertEqual((evaluacion.compromisos_previos, evaluacion.disponible_despues),
                         (Decimal('200'), Decimal('9700')))

    def test_detalle_no_excluye_compromiso_de_otra_cotizacion(self):
        self.ordenar()
        # Una selección inconsistente no debe ocultar la exposición de la orden.
        self.item.cotizaciones.update(seleccionada=False)
        CotizacionCompraDepartamental.objects.create(
            item=self.item, proveedor=self.proveedor, cantidad_ofertada=1,
            costo_unitario=100, seleccionada=True,
        )
        detalle = self.client.get(reverse('compras:departamental_detalle', args=[self.solicitud.pk]))
        visible = list(detalle.context['solicitud'].items.all())[0]
        self.assertEqual(visible.evaluacion_presupuesto.disponible_antes, Decimal('9800'))
        self.assertEqual(visible.evaluacion_presupuesto.disponible_despues, Decimal('9700'))

    def test_reembolso_pendiente_sobrevive_cierre_y_filtro_estado_sin_mezclar_kpi(self):
        self.quote.costo_unitario = 500
        self.quote.save(update_fields=['costo_unitario'])
        self.assertEqual(self.comprar(importe_final='1000').status_code, 200)
        anterior = self.item.intento_vigente
        self.cancelar(anterior, pagado=True)
        anterior.refresh_from_db()
        registrar_reembolso_compra(anterior, version=anterior.version, fecha=timezone.localdate(),
                                  importe=Decimal('250'), actor=self.user)
        nuevo = self.reemplazar()
        self.assertEqual(self.recibir(nuevo, '2').status_code, 200)
        confirmar_recepcion_departamental(self.item, conforme=True, actor=self.user)
        self.solicitud.refresh_from_db()
        self.assertEqual(self.solicitud.estado, 'COMPLETADA')
        params = {'periodo': '2026-09', 'area': self.area.pk, 'estado': 'POR_COTIZAR'}
        url = reverse('compras:departamental_bandeja')
        response = self.client.get(url, params)
        self.assertEqual([i.pk for i in response.context['items']], [self.item.pk])
        total = response.context['resumen']
        self.assertEqual((total['reembolso_pendiente'], total['reembolsado']), (Decimal('750'), Decimal('250')))
        for campo in ('articulos', 'solicitudes', 'solicitado', 'cotizado', 'comprometido', 'sin_precio'):
            self.assertEqual(total[campo], 0, campo)
        self.assertContains(response, 'Solo seguimiento de reembolso')
        self.assertContains(response, '$750.00')
        wb = load_workbook(BytesIO(self.client.get(url, {**params, 'exportar': 'xlsx'}).content))
        for hoja in wb:
            headers = [c.value for c in hoja[8]]
            self.assertEqual(hoja.cell(9, headers.index('Reembolso pendiente') + 1).value, 750)
            self.assertEqual(hoja.cell(9, headers.index('Reembolsado') + 1).value, 250)
        self.assertEqual(construir_resumen_departamental({'periodo': '2026-08'})['items'], [])
        from reportes.models import AreaPresupuesto
        otra_area = AreaPresupuesto.objects.create(nombre='Otra área', codigo='otra-cruces')
        self.assertEqual(construir_resumen_departamental({'area': otra_area.pk})['items'], [])
        anterior.refresh_from_db()
        registrar_reembolso_compra(anterior, version=anterior.version, fecha=timezone.localdate(),
                                  importe=Decimal('750'), actor=self.user)
        # Es bandeja de pendientes, no libro histórico: el cierre financiero permite salir.
        self.assertEqual(construir_resumen_departamental({})['items'], [])
        self.assertEqual(ReembolsoCompraDepartamental.objects.filter(intento=anterior).count(), 2)
        detalle = self.client.get(reverse('compras:departamental_detalle', args=[self.solicitud.pk]))
        self.assertContains(detalle, 'Reembolso completado')
        self.assertContains(detalle, '$1,000.00')
