"""Consulta de pendientes compartida por la bandeja y su exportación."""
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from urllib.parse import urlencode

from django import forms
from django.db.models import Prefetch
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportes.models import AreaPresupuesto
from .models import CotizacionCompraDepartamental, ItemCompraDepartamental, SolicitudCompraDepartamental

TERMINALES = (
    ItemCompraDepartamental.ESTADO_RECIBIDO_CONFORME,
    ItemCompraDepartamental.ESTADO_RECHAZADO,
    ItemCompraDepartamental.ESTADO_CANCELADO,
)
ESTADOS = [(value, label) for value, label in ItemCompraDepartamental.ESTADO_CHOICES if value not in TERMINALES]
ALCANCE = 'Solo artículos pendientes de solicitudes enviadas. Excluye borradores, canceladas, completadas, rechazados y recibidos conforme.'
ETAPAS = 'Importes en MXN. Solicitado, cotizado y comprometido son etapas distintas: no se suman entre sí.'


class FiltrosResumenDepartamental(forms.Form):
    periodo = forms.DateField(
        label='Mes planeado', required=False, input_formats=['%Y-%m'],
        widget=forms.DateInput(format='%Y-%m', attrs={'type': 'month'}),
    )
    area = forms.ModelChoiceField(
        label='Departamento', required=False, empty_label='Todos los departamentos',
        queryset=AreaPresupuesto.objects.order_by('nombre'),
    )
    estado = forms.ChoiceField(label='Estado del artículo', required=False, choices=[('', 'Todos los pendientes'), *ESTADOS],
                              widget=forms.Select(attrs={'data-native-select': 'true'}))


def _importe(value):
    return value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _totales():
    return dict(solicitudes=set(), articulos=0, solicitado=Decimal('0'), cotizado=Decimal('0'),
                comprometido=Decimal('0'), sin_estimacion=0, sin_cotizacion=0, sin_precio=0)


def construir_resumen_departamental(params):
    filtros = FiltrosResumenDepartamental(params)
    valido = filtros.is_valid()
    items = ItemCompraDepartamental.objects.select_related(
        'solicitud__area', 'solicitud__solicitante', 'solicitud__comprador_asignado', 'compromiso',
    ).exclude(estado__in=TERMINALES).exclude(solicitud__estado__in=[
        SolicitudCompraDepartamental.ESTADO_BORRADOR,
        SolicitudCompraDepartamental.ESTADO_CANCELADA,
        SolicitudCompraDepartamental.ESTADO_COMPLETADA,
    ]).prefetch_related(Prefetch(
        'cotizaciones', queryset=CotizacionCompraDepartamental.objects.filter(seleccionada=True).order_by('id'),
        to_attr='cotizaciones_seleccionadas',
    )).order_by('solicitud__area__nombre', 'solicitud_id', 'id')
    query = {}
    if not valido:
        items = items.none()
    else:
        for campo, lookup in [('periodo', 'solicitud__periodo'), ('area', 'solicitud__area'), ('estado', 'estado')]:
            valor = filtros.cleaned_data[campo]
            if valor:
                items = items.filter(**{lookup: valor})
                query[campo] = valor.strftime('%Y-%m') if campo == 'periodo' else (valor.pk if campo == 'area' else valor)
    items = list(items)
    total = _totales()
    departamentos = {}
    for item in items:
        area = item.solicitud.area
        grupo = departamentos.setdefault(area.pk, {**_totales(), 'nombre': area.nombre, 'id': area.pk})
        estimado = item.subtotal_estimado
        quote = next(iter(item.cotizaciones_seleccionadas), None)
        compromiso = getattr(item, 'compromiso', None)
        item.resumen_estimado = _importe(estimado) if estimado is not None else None
        item.resumen_cotizado = _importe(quote.total_adquisicion) if quote else None
        item.resumen_comprometido = compromiso.monto if compromiso and compromiso.activo and compromiso.formalizado_en else Decimal('0')
        item.resumen_sin_precio = estimado is None and quote is None
        for destino in (total, grupo):
            destino['solicitudes'].add(item.solicitud_id)
            destino['articulos'] += 1
            destino['solicitado'] += item.resumen_estimado or Decimal('0')
            destino['cotizado'] += item.resumen_cotizado or Decimal('0')
            destino['comprometido'] += item.resumen_comprometido
            destino['sin_estimacion'] += estimado is None
            destino['sin_cotizacion'] += quote is None
            destino['sin_precio'] += item.resumen_sin_precio
    for grupo in [total, *departamentos.values()]:
        grupo['solicitudes'] = len(grupo['solicitudes'])
    base = reverse('compras:departamental_bandeja')
    for grupo in departamentos.values():
        grupo['url'] = base + '?' + urlencode({**query, 'area': grupo['id']}) + '#articulos'
    return {
        'filtros': filtros, 'items': items, 'resumen': total,
        'departamentos': list(departamentos.values()), 'alcance_resumen': ALCANCE,
        'etapas_resumen': ETAPAS, 'exportar_url': base + '?' + urlencode({**query, 'exportar': 'xlsx'}),
    }


def exportar_resumen_departamental(contexto):
    """Exporta los mismos renglones calculados, sin fórmulas desde texto capturado."""
    filtros = contexto['filtros']
    datos = filtros.cleaned_data
    filtros_texto = (
        f"Mes planeado: {datos['periodo']:%Y-%m}" if datos['periodo'] else 'Mes planeado: todos'
    ) + f" · Departamento: {datos['area'].nombre if datos['area'] else 'todos'}" + f" · Estado: {dict(ESTADOS).get(datos['estado'], 'todos los pendientes')}"
    wb = Workbook()
    wb.active.title = 'Resumen'
    detalle = wb.create_sheet('Artículos')
    fecha = timezone.localtime().strftime('%Y-%m-%d %H:%M %Z')
    total = contexto['resumen']
    cobertura = ('Total estimado parcial. ' if total['sin_estimacion'] else '') + f"Artículos sin estimación: {total['sin_estimacion']}. Sin cotización seleccionada: {total['sin_cotizacion']}. Sin ningún precio: {total['sin_precio']}."
    for ws in wb:
        for text in ('Compras departamentales · pendientes', filtros_texto, ALCANCE, ETAPAS,
                     f'Generado: {fecha}', cobertura):
            ws.append([text])
        ws.append([])
    resumen_headers = ['Departamento', 'Solicitudes', 'Artículos', 'Sin precio', 'Solicitado estimado', 'Cotizado', 'Comprometido', 'Sin estimación', 'Sin cotización']
    wb.active.append(resumen_headers)
    for grupo in [*contexto['departamentos'], {'nombre': 'Total', **total}]:
        wb.active.append([grupo['nombre'], *[grupo[key] for key in (
            'solicitudes', 'articulos', 'sin_precio', 'solicitado', 'cotizado', 'comprometido', 'sin_estimacion', 'sin_cotizacion',
        )]])
    detalle.append(['Solicitud', 'Departamento', 'Mes planeado', 'Artículo', 'Cantidad', 'Unidad',
                    'Estado', 'Costo unitario estimado', 'Solicitado estimado', 'Cotizado', 'Comprometido', 'Sin precio'])
    for item in contexto['items']:
        detalle.append([item.solicitud.folio, item.solicitud.area.nombre, item.solicitud.periodo.strftime('%Y-%m'),
                        item.descripcion, item.cantidad, item.unidad, item.get_estado_display(), item.costo_unitario_estimado,
                        item.resumen_estimado, item.resumen_cotizado, item.resumen_comprometido,
                        'Sí' if item.resumen_sin_precio else 'No'])
    for ws in wb:
        ws.freeze_panes = 'B9'
        ws.auto_filter.ref = f'A8:{get_column_letter(ws.max_column)}{ws.max_row - (1 if ws.title == "Resumen" else 0)}'
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if cell.row == 8:
                    cell.fill = PatternFill('solid', fgColor='8B2252')
                    cell.font = Font(color='FFFFFF', bold=True)
                if cell.row >= 8:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                money_cols = (5, 6, 7) if ws.title == 'Resumen' else (8, 9, 10, 11)
                if cell.row > 8 and cell.column in money_cols:
                    cell.number_format = '"$"#,##0.00'
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 23
        if ws.title == 'Artículos':
            ws.column_dimensions['D'].width = 52
            for cell in ws['E'][8:]:
                cell.number_format = '0.###'
        else:
            ws.column_dimensions['A'].width = 32
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True, color='8B2252')
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '8:8'
    output = BytesIO()
    wb.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="compras_departamentales_resumen.xlsx"'
    response['Cache-Control'] = 'private, no-store'
    return response
