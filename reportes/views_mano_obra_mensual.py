from datetime import timedelta
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from reportes.models import CierreManoObraMensual
from reportes.services_mano_obra_mensual import aprobar, calcular, limites, puede_revisar


@login_required
@never_cache
@require_http_methods(['GET', 'POST'])
def mensual(request):
    if not puede_revisar(request.user):
        raise PermissionDenied
    previo = timezone.localdate().replace(day=1) - timedelta(days=1)
    periodo = request.POST.get('periodo') if request.method == 'POST' else request.GET.get('periodo', previo.strftime('%Y-%m'))
    try:
        inicio, _ = limites(periodo)
    except ValidationError as exc:
        return HttpResponse(exc.messages[0], status=400)
    if request.method == 'POST':
        if request.POST.get('confirmar_fuentes') != 'si':
            messages.error(request, 'Confirma la cobertura del mes y la clasificación histórica del personal.')
        else:
            try:
                cierre = aprobar(periodo, request.POST.get('fingerprint', ''), request.user)
                messages.success(request, f'Cierre aprobado. Revisión {cierre.revision}. Disponible en Finanzas.')
            except ValidationError as exc:
                messages.error(request, ' '.join(exc.messages))
        return redirect(f"{reverse('reportes:mano_obra_mensual')}?periodo={periodo}")
    datos = calcular(periodo)
    if request.GET.get('exportar') == 'xlsx':
        return exportar(datos)
    historial = CierreManoObraMensual.objects.filter(periodo=inicio).select_related('aprobado_por').order_by('-revision')
    return render(request, 'reportes/mano_obra_mensual.html', {'mo': datos, 'historial': historial, 'cierre': historial.first()})


def exportar(datos):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Simulacion'
    sheet.append(['SIMULACIÓN — ' + datos['periodo'], 'No aplicada a los análisis'])
    sheet.append(['Fondo', datos['fondo'], 'Asignado', datos['asignado']])
    sheet.append(['Producto', 'Código', 'Elaborado', 'Materia prima por unidad', 'Elaboración por unidad', 'Supervisión por unidad', 'Despacho por unidad', 'Mano de obra por unidad', 'Fabricación por unidad', 'Con despacho por unidad', 'Pendientes', 'Servicios por unidad', 'Fabricación con servicios', 'Con servicios y despacho'])
    for r in datos['productos']:
        componentes = [c['unitario'] for c in r['componentes']] or [None] * 3
        sheet.append([r['nombre'], r['codigo'], r['cantidad'], r['mp'], *componentes, r['mo'], r['fabricacion'], r['con_despacho'], ' '.join(r['errores']), r['servicios'], r['fabricacion_servicios'], r['con_despacho_servicios']])
    servicios = wb.create_sheet('Servicios')
    servicios.append(['Servicio', 'Real de Producción'])
    for s in datos['servicios']['fondos']:
        servicios.append([s['nombre'], s['monto']])
    servicios.append(['Total', datos['servicios']['fondo']])
    servicios.append(['Asignado', datos['servicios']['asignado']])
    servicios.append(['Pendiente', datos['servicios']['pendiente']])
    servicios.append(['Producto', 'Servicio', 'Asignado mensual', 'Por unidad'])
    for r in datos['productos']:
        for s in r['servicios_detalle']:
            servicios.append([r['nombre'], s['nombre'], s['monto'], s['unitario']])
    personal = wb.create_sheet('Personal')
    personal.append(['Nombre', 'Puesto', 'Componente', 'Percepciones'])
    for p in datos['personas']:
        personal.append([p['nombre'], p['puesto'], p['componente'], p['monto']])
    fuente = wb.create_sheet('Fuentes y conciliacion')
    for row in [['Periodo', datos['periodo']], ['Huella de fuentes', datos['fingerprint']], ['Fondo', datos['fondo']], ['Asignado', datos['asignado']], ['Redondeo', datos['redondeo']], ['Pendiente nómina', datos['pendiente_nomina']], ['Pendiente reparto', datos['pendiente_reparto']]]:
        fuente.append(row)
    for nota in datos['pendientes'] + datos['notas']:
        fuente.append([nota])
    for ws in wb:
        ws.freeze_panes = 'C4' if ws == sheet else 'A2'
        for cells in ws:
            for cell in cells:
                # Los nombres vienen de fuentes externas: nunca exportarlos como fórmulas.
                if cell.data_type == 'f':
                    cell.data_type = 's'
                if isinstance(cell.value, (int, float)) or cell.data_type == 'n':
                    cell.number_format = '#,##0.00'
        for cell in ws[3 if ws == sheet else 1]:
            cell.fill = PatternFill('solid', fgColor='8B2252')
            cell.font = Font(color='FFFFFF', bold=True)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 28
        ws.column_dimensions['A'].width = 48
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="mano-obra-{datos["periodo"]}.xlsx"'
    return response
