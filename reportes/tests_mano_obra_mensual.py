from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.db import DatabaseError, transaction
from django.test import TestCase, Client
from django.urls import reverse
from openpyxl import load_workbook

from maestros.models import UnidadMedida
from pos_bridge.models import PointBranch, PointProductionLine
from recetas.models import Receta, RecetaCodigoPointAlias
from reportes.models import CierreManoObraMensual, RecetaCostoHistoricoMensual, AreaPresupuesto, RubroPresupuesto, LineaPresupuestoMensual
from reportes.services_mano_obra_mensual import calcular, leer_fuentes, aprobar, analisis_aprobado
from rrhh.models import Empleado, NominaLinea, NominaPeriodo


class ManoObraMensualTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('mo-admin', 'mo@example.invalid', 'test-only-password')
        cls.periodo = NominaPeriodo.objects.create(folio='MO-2026-09', fecha_inicio=date(2026, 9, 1), fecha_fin=date(2026, 9, 30), estatus='PAGADA')
        cls.persona = Empleado.objects.create(codigo='MO-1', nombre='Responsable prueba', departamento='PRODUCCION', departamento_origen='PRODUCCION', puesto='Jefe de Producción', puesto_operativo='')
        cls.nomina = NominaLinea.objects.create(periodo=cls.periodo, empleado=cls.persona, departamento_snapshot='PRODUCCION', total_percepciones=Decimal('100.01'))
        NominaLinea.objects.filter(pk=cls.nomina.pk).update(total_percepciones=Decimal('100.01'))
        cls.unidad = UnidadMedida.objects.create(codigo='PZA', nombre='Pieza')
        cls.branch = PointBranch.objects.create(external_id='MO-LOCAL', name='Producción prueba')
        cls.receta = Receta.objects.create(nombre='Pastel prueba', codigo_point='MO-P1', tipo='PRODUCTO_FINAL', hash_contenido='mo-p1', rendimiento_cantidad=1, rendimiento_unidad=cls.unidad)
        cls.linea = PointProductionLine.objects.create(branch=cls.branch, receta=cls.receta, production_external_id='MO-1', detail_external_id='MO-D1', source_hash='mo-1', production_date=date(2026, 9, 15), item_name='Pastel prueba', item_code='MO-P1', unit='PZA', produced_quantity=10)
        cls.costo = RecetaCostoHistoricoMensual.objects.create(periodo=date(2026, 9, 1), receta=cls.receta, costo_total=20, costo_por_unidad_rendimiento=20, coverage_pct=100)
        area, _ = AreaPresupuesto.objects.get_or_create(codigo='produccion', defaults={'nombre': 'Producción'})
        for nombre in ('Agua potable', 'Energía eléctrica', 'Gas'):
            rubro = RubroPresupuesto.objects.create(area=area, concepto=nombre, tipo='EGRESO')
            LineaPresupuestoMensual.objects.create(rubro=rubro, periodo=date(2026, 9, 1),
                monto_real=0, fuente_real='AUTO:GASTO_OPERATIVO')

    def test_servicios_conservan_centavos_y_reparten_por_mp(self):
        fuentes = leer_fuentes('2026-09')
        for s in fuentes['servicios']:
            s['monto_real'] = '1.01'
        # Dos productos con igual base: no perder centavos por redondeo.
        produccion = fuentes['produccion'][0]
        fuentes['produccion'].append({**produccion, 'id': 999, 'receta_id': 999,
            'production_external_id': 'otra', 'detail_external_id': 'otra',
            'receta__nombre': 'Otro producto'})
        fuentes['costos'].append({**fuentes['costos'][0], 'receta_id': 999})
        data = calcular('2026-09', fuentes)
        self.assertEqual(data['servicios']['fondo'], Decimal('3.03'))
        self.assertEqual(data['servicios']['asignado'], Decimal('3.03'))
        self.assertEqual(data['servicios']['pendiente'], 0)
        self.assertEqual(sorted(p['servicios_asignados'] for p in data['productos']), [Decimal('1.50'), Decimal('1.53')])
        for p in data['productos']:
            self.assertEqual(p['con_despacho_servicios'], p['con_despacho'] + p['servicios'])

    def test_servicio_faltante_no_se_convierte_en_cero(self):
        LineaPresupuestoMensual.objects.filter(rubro__concepto='Gas').update(monto_real=None)
        data = calcular('2026-09')
        self.assertIsNone(data['servicios']['fondo'])
        self.assertIsNone(data['productos'][0]['con_despacho_servicios'])
        self.assertIsNotNone(data['productos'][0]['mo'])
        with self.assertRaises(ValidationError):
            aprobar('2026-09', data['fingerprint'], self.user)

    def test_servicios_no_duplican_version_ni_area_y_cambio_invalida_huella(self):
        linea = LineaPresupuestoMensual.objects.get(rubro__concepto='Gas')
        LineaPresupuestoMensual.objects.create(rubro=linea.rubro, periodo=linea.periodo,
            version='REVISADO', monto_real=900, fuente_real='AUTO:GASTO_OPERATIVO')
        antes = calcular('2026-09')
        self.assertEqual(antes['servicios']['fondo'], 0)
        linea.monto_real = 10
        linea.save()
        with self.assertRaises(ValidationError):
            aprobar('2026-09', antes['fingerprint'], self.user)
        data = calcular('2026-09')
        self.assertEqual(data['servicios']['fondo'], 10)
        aprobar('2026-09', data['fingerprint'], self.user)
        self.assertEqual(analisis_aprobado('2026-09')['filas'][0]['costo_comparado'], Decimal('31.001'))

    def test_servicios_rechazan_ambiguedad_negativos_y_fuente_no_real(self):
        for cambio in ({'monto_real': '-1'}, {'fuente_real': 'PRESUPUESTO'}):
            fuentes = leer_fuentes('2026-09')
            fuentes['servicios'][0].update(cambio)
            self.assertIsNone(calcular('2026-09', fuentes)['servicios']['fondo'])
        fuentes = leer_fuentes('2026-09')
        fuentes['servicios'].append(fuentes['servicios'][0].copy())
        self.assertIsNone(calcular('2026-09', fuentes)['servicios']['fondo'])

    def test_cierre_previo_no_gana_servicios_retroactivos(self):
        from django.core.serializers.json import DjangoJSONEncoder
        import json
        data = calcular('2026-09')
        data.pop('servicios')
        for r in data['productos']:
            r.pop('con_despacho_servicios')
        CierreManoObraMensual.objects.create(periodo=date(2026, 9, 1), revision=1,
            fingerprint='historico', aprobado_por=self.user,
            resultado=json.loads(json.dumps(data, cls=DjangoJSONEncoder)))
        self.assertEqual(analisis_aprobado('2026-09')['filas'][0]['costo_comparado'], Decimal('30.001'))

    def test_fuentes_y_cargo_alternativo(self):
        fuentes = leer_fuentes('2026-09')
        self.assertEqual(fuentes['nomina'][0]['empleado__puesto'], 'Jefe de Producción')
        data = calcular('2026-09')
        self.assertEqual(data['pendientes'], [])
        self.assertEqual(data['fondo'], Decimal('100.01'))
        self.assertEqual(data['productos'][0]['mo'], Decimal('10.001'))
        self.assertEqual(data['personas'][0]['componente'], 'Supervisión y control')

    def test_inactivo_conserva_nomina(self):
        Empleado.objects.filter(pk=self.persona.pk).update(activo=False)
        self.assertEqual(calcular('2026-09')['fondo'], Decimal('100.01'))

    def test_despacho_separado_fabricacion(self):
        Empleado.objects.filter(pk=self.persona.pk).update(puesto_operativo='ENVIO_SUCURSAL')
        r = calcular('2026-09')['productos'][0]
        self.assertEqual(r['fabricacion'], Decimal('20'))
        self.assertEqual(r['con_despacho'], Decimal('30.001'))

    def test_faltante_costo_no_reparte_a_los_demas(self):
        otra = Receta.objects.create(nombre='Sin costo', codigo_point='OTRA', tipo='PRODUCTO_FINAL', hash_contenido='otra', rendimiento_cantidad=1, rendimiento_unidad=self.unidad)
        PointProductionLine.objects.create(branch=self.branch, receta=otra, production_external_id='MO-2', detail_external_id='MO-D2', source_hash='mo-2', production_date=date(2026, 9, 15), item_name='Sin costo', item_code='OTRA', unit='PZA', produced_quantity=5)
        data = calcular('2026-09')
        self.assertTrue(all(p['mo'] is None for p in data['productos']))
        self.assertEqual(data['pendiente_reparto'], Decimal('100.01'))

    def test_receta_por_pieza_no_exige_rendimiento_batida(self):
        Receta.objects.filter(pk=self.receta.pk).update(rendimiento_cantidad=None, rendimiento_unidad=None)
        RecetaCostoHistoricoMensual.objects.filter(pk=self.costo.pk).update(costo_por_unidad_rendimiento=None)
        data = calcular('2026-09')
        self.assertEqual(data['pendientes'], [])
        self.assertEqual(data['productos'][0]['mp'], Decimal('20'))
        self.assertEqual(data['productos'][0]['mo'], Decimal('10.001'))
        aprobar('2026-09', data['fingerprint'], self.user)
        self.receta.refresh_from_db()
        self.assertIsNone(self.receta.rendimiento_cantidad)

    def test_rendimiento_explicito_invalido_bloquea(self):
        for rendimiento in (0, -1):
            Receta.objects.filter(pk=self.receta.pk).update(rendimiento_cantidad=rendimiento)
            data = calcular('2026-09')
            with self.assertRaises(ValidationError):
                aprobar('2026-09', data['fingerprint'], self.user)

    def test_alias_existente_resuelve_preparacion_sin_modificar_fuente(self):
        receta = Receta.objects.create(nombre='Pan horneado', tipo='PREPARACION', codigo_point='PMH028', hash_contenido='pmh')
        alias = RecetaCodigoPointAlias.objects.create(receta=receta, codigo_point='PMH028')
        linea = PointProductionLine.objects.create(branch=self.branch, production_external_id='MO-2', detail_external_id='MO-D2', source_hash='mo-2', production_date=date(2026, 9, 15), item_name='Pan horneado', item_code='PMH028', unit='PZA', produced_quantity=26, is_insumo=True)
        data = calcular('2026-09')
        self.assertEqual(data['pendientes'], [])
        self.assertEqual(data['preparaciones'], 1)
        self.assertEqual(data['unidades'], 10)
        linea.refresh_from_db()
        self.assertIsNone(linea.receta_id)
        alias.activo = False
        alias.save()
        self.assertEqual(calcular('2026-09')['sin_vinculo'], 1)

    def test_costo_lote_se_divide_por_rendimiento(self):
        Receta.objects.filter(pk=self.receta.pk).update(rendimiento_cantidad=5)
        RecetaCostoHistoricoMensual.objects.filter(pk=self.costo.pk).update(costo_por_unidad_rendimiento=None)
        self.assertEqual(calcular('2026-09')['productos'][0]['mp'], Decimal('4'))

    def test_preparaciones_no_reciben_segundo_reparto(self):
        Receta.objects.filter(pk=self.receta.pk).update(tipo='PREPARACION')
        data = calcular('2026-09')
        self.assertEqual(data['productos'], [])
        self.assertEqual(data['preparaciones'], 1)
        self.assertEqual(data['asignado'], 0)

    def test_periodo_incompleto_y_sin_nomina_no_es_cero_confiable(self):
        NominaPeriodo.objects.filter(pk=self.periodo.pk).update(fecha_inicio=date(2026, 9, 16))
        self.assertTrue(any('todos los días' in x for x in calcular('2026-09')['pendientes']))
        self.assertTrue(calcular('2026-08')['pendientes'])

    def test_nomina_ausente_o_sin_clasificar_no_es_mano_obra_cero(self):
        NominaPeriodo.objects.filter(pk=self.periodo.pk).update(estatus='BORRADOR')
        self.assertIsNone(calcular('2026-09')['productos'][0]['mo'])
        NominaPeriodo.objects.filter(pk=self.periodo.pk).update(estatus='PAGADA')
        Empleado.objects.filter(pk=self.persona.pk).update(puesto='Sin clasificación')
        self.assertIsNone(calcular('2026-09')['productos'][0]['mo'])

    def test_rechaza_fuente_cambiada_y_preserva_revision(self):
        data = calcular('2026-09')
        first = aprobar('2026-09', data['fingerprint'], self.user)
        self.assertEqual(aprobar('2026-09', data['fingerprint'], self.user).pk, first.pk)
        NominaLinea.objects.filter(pk=self.nomina.pk).update(total_percepciones=200)
        with self.assertRaises(ValidationError):
            aprobar('2026-09', data['fingerprint'], self.user)
        newer = aprobar('2026-09', calcular('2026-09')['fingerprint'], self.user)
        self.assertEqual(newer.revision, 2)
        self.assertEqual(analisis_aprobado('2026-09')['cierre'].pk, newer.pk)
        first.refresh_from_db()
        self.assertEqual(first.resultado['fondo'], '100.01')

    def test_cierre_inmutable_en_bd(self):
        cierre = aprobar('2026-09', calcular('2026-09')['fingerprint'], self.user)
        with self.assertRaises(DatabaseError), transaction.atomic():
            CierreManoObraMensual.objects.filter(pk=cierre.pk).update(fingerprint='otro')
        with self.assertRaises(DatabaseError), transaction.atomic():
            CierreManoObraMensual.objects.filter(pk=cierre.pk).delete()

    def test_analisis_sin_cierre_y_sin_venta(self):
        self.assertIsNone(analisis_aprobado('2026-09')['cierre'])
        aprobar('2026-09', calcular('2026-09')['fingerprint'], self.user)
        self.assertIsNone(analisis_aprobado('2026-09')['filas'][0]['margen'])

    def test_vista_get_no_escribe_y_exporta(self):
        self.client.force_login(self.user)
        url = reverse('reportes:mano_obra_mensual')
        result = self.client.get(url, {'periodo': '2026-09'})
        self.assertEqual(result.status_code, 200)
        self.assertContains(result, 'Pastel prueba')
        self.assertFalse(CierreManoObraMensual.objects.exists())
        result = self.client.get(url, {'periodo': '2026-09', 'exportar': 'xlsx'})
        wb = load_workbook(BytesIO(result.content))
        self.assertEqual(wb['Simulacion']['H4'].value, 10.001)
        self.assertEqual(self.client.get(url, {'periodo': 'invalido'}).status_code, 400)

    def test_permisos_csrf_confirmacion(self):
        url = reverse('reportes:mano_obra_mensual')
        otro = User.objects.create_user('lector', password='test-password')
        self.client.force_login(otro)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.post(url, {'periodo': '2026-09'}).status_code, 403)
        self.client.force_login(self.user)
        self.client.post(url, {'periodo': '2026-09', 'fingerprint': calcular('2026-09')['fingerprint']})
        self.assertFalse(CierreManoObraMensual.objects.exists())
        csrf = Client(enforce_csrf_checks=True)
        csrf.force_login(self.user)
        response = csrf.post(url, {'periodo': '2026-09', 'confirmar_fuentes': 'si', 'fingerprint': calcular('2026-09')['fingerprint']})
        self.assertRedirects(response, reverse('login'), fetch_redirect_response=False)
        self.assertFalse(CierreManoObraMensual.objects.exists())

    def test_exportacion_no_inyecta_formula(self):
        Receta.objects.filter(pk=self.receta.pk).update(nombre='=1+1')
        self.client.force_login(self.user)
        response = self.client.get(reverse('reportes:mano_obra_mensual'), {'periodo': '2026-09', 'exportar': 'xlsx'})
        cell = load_workbook(BytesIO(response.content))['Simulacion']['A4']
        self.assertEqual(cell.value, '=1+1')
        self.assertEqual(cell.data_type, 's')
